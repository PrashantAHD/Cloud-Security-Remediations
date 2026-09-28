"""Import local-account execution evidence; never execute commands or contact a cloud."""

import argparse
import csv
import json
import re
import sys
from collections import Counter
from collections.abc import Sequence
from copy import copy
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from textwrap import wrap
from uuid import uuid4
from zipfile import BadZipFile

from defusedxml.ElementTree import ParseError
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from remediation.csv_report import REPOSITORY_ROOT, _output_path, _publish, _validate_text

REQUIRED = ("Account", "Resource Name", "Region", "Remediation Status", "Local User Action")
EXTRA = ("Blocker", "Next Action")
DATA = ("AWS Data", "Azure Data")
PENDING = ("Pending Review", "Pending Remediation")
COLORS = {
    "Accepted risk": "FFC6EFCE", "Remediated": "FFFFC7CE",
    "Pending review": "FFFFEB9C", "Pending remediation": "FFFFEB9C", "Blocked": "FFFFEB9C",
}
SUCCESS = "ACCOUNT_REMOVAL_VERIFIED_HOME_PATHS_PRESERVED_SSH_CONFIG_UNCHANGED"
NO_CHANGE = "NO_CHANGES_ALL_TARGETS_ALREADY_ABSENT"
STOPPED = "STOPPED_REVIEW_REQUIRED"
DISCLAIMER = "Application-owner validation and fresh Wiz reassessment remain required."
USER = r"[a-z_][a-z0-9_.-]*\$?"
MARKERS = re.compile(
    r"(PYTHON_INTERPRETER|VM|EXECUTION|SYSTEMD_TEMPLATE_DEFINITIONS|ACCOUNT|PRECHECKS|"
    r"BACKUP_PATH|REMOVING|REMOVED_AND_VERIFIED|REMOVED_THIS_RUN|STOP|"
    r"REMOVAL_COMMANDS_SUCCEEDED|RESULT)="
)


@dataclass(frozen=True)
class ExecutionResult:
    vm: str
    outcome: str
    verified: tuple[str, ...]
    succeeded: tuple[str, ...]
    absent: tuple[str, ...]
    backup: str
    reason: str
    accounts: tuple[str, ...]


@dataclass(frozen=True)
class UpdateResult:
    imported: int
    duplicates: int
    backup_path: Path | None
    memory_path: Path


def _users(value: str, delimiter: str = ",") -> tuple[str, ...]:
    users = tuple(part.strip() for part in value.split(delimiter))
    if not users or any(not re.fullmatch(USER, user) for user in users):
        raise ValueError("Invalid or missing local-account names in evidence or approved scope.")
    if len(set(users)) != len(users):
        raise ValueError("Duplicate local-account names in evidence or scope.")
    return users


def _text(value: str) -> str:
    _validate_text(value)
    if len(value) > 32767:
        raise ValueError("Evidence exceeds the Excel cell text limit.")
    return value


def parse_result(text: str) -> ExecutionResult:
    """Parse the supported marker protocol, including newline-collapsed portal output.

    These are operator-supplied statements, not authenticated guest telemetry.
    Unknown protocols, incomplete output and contradictory success fail closed.
    """
    if len(text) > 1_000_000:
        raise ValueError("Run Command output exceeds the supported size.")
    _validate_text(text)
    text = text.strip()
    if "[stderr]" in text:
        text, stderr = text.split("[stderr]", 1)
        if stderr.strip():
            raise ValueError("Nonempty stderr requires manual review; use a blocker entry.")
    if "[stdout]" in text:
        prefix, text = text.split("[stdout]", 1)
        if prefix.strip() not in ("", "Enable succeeded:"):
            raise ValueError("Unsupported Run Command envelope.")
    text = text.replace(
        "WARNING: Earlier Wiz export marked this VM inactive; confirm current ownership/state.", "",
    )
    matches = list(MARKERS.finditer(text))
    if not matches or text[:matches[0].start()].strip():
        raise ValueError("No supported Run Command evidence found.")
    fields: dict[str, list[str]] = {}
    ordered: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        value = text[match.end():end].strip()
        fields.setdefault(match[1], []).append(value)
        ordered.append((match[1], value))

    def one(name: str, required: bool = True) -> str:
        values = fields.get(name, [])
        if len(values) > 1 or (required and not values):
            raise ValueError(f"Missing or repeated {name} marker; do not merge executions.")
        return values[0] if values else ""

    vm = one("VM")
    one("PYTHON_INTERPRETER", required=False)
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", vm):
        raise ValueError("Invalid VM marker.")
    if one("EXECUTION") != "SINGLE_RUN":
        raise ValueError("Unsupported execution protocol.")
    terminal = one("RESULT")
    if ordered[-1][0] != "RESULT":
        raise ValueError("Result marker must terminate the execution.")
    if terminal.startswith(SUCCESS):
        if terminal[len(SUCCESS):].strip() not in ("", DISCLAIMER):
            raise ValueError("Unexpected text after success result.")
        outcome = "verified"
    elif terminal == NO_CHANGE:
        outcome = "already_absent"
    elif terminal == STOPPED or terminal.startswith(STOPPED + ";"):
        outcome = "stopped"
    else:
        raise ValueError("Unknown or incomplete execution result.")
    precheck = one("PRECHECKS", required=outcome != "stopped")
    if precheck and precheck != "PASSED; proceeding automatically with approved removals.":
        raise ValueError("Unrecognized precheck result.")
    reason = one("STOP", required=outcome == "stopped")
    backup = one("BACKUP_PATH", required=False)
    if backup and not re.fullmatch(
        r"/var/lib/cloud-security-remediation/account-removal-[A-Za-z0-9_-]+", backup
    ):
        raise ValueError("Unexpected backup path.")
    accounts, absent = [], []
    for value in fields.get("ACCOUNT", []):
        missing = re.fullmatch(
            rf"({USER}) ALREADY_ABSENT \(not proof of who removed it\)", value,
        )
        present = re.fullmatch(
            rf"({USER}) UID=\d+ GID=\d+ HOME=/[^\s]* SHELL=/[^\s]+", value,
        )
        if not (missing or present):
            raise ValueError("Invalid account observation.")
        name = (missing or present)[1]
        if name in accounts:
            raise ValueError("Repeated account observation.")
        accounts.append(name)
        if missing:
            absent.append(name)
    removing = tuple(fields.get("REMOVING", []))
    verified = tuple(fields.get("REMOVED_AND_VERIFIED", []))
    for names in (removing, verified):
        if names:
            _users(",".join(names))
    if not set(verified) <= set(removing) <= set(accounts) - set(absent):
        raise ValueError("Removal markers contradict account observations.")
    if verified != removing[:len(verified)]:
        raise ValueError("Verified accounts must follow the attempted removal order.")
    for name in verified:
        if ordered.index(("REMOVING", name)) > ordered.index(("REMOVED_AND_VERIFIED", name)):
            raise ValueError("Verification precedes removal.")
    succeeded_text = one("REMOVAL_COMMANDS_SUCCEEDED", required=outcome == "stopped")
    succeeded = (
        _users(succeeded_text) if succeeded_text and succeeded_text != "none" else ()
    )
    if not set(verified) <= set(succeeded if outcome == "stopped" else removing):
        raise ValueError("Verification conflicts with command success list.")
    if not set(succeeded) <= set(removing):
        raise ValueError("Command success list contains an unattempted account.")
    if removing and (not backup or not precheck):
        raise ValueError("Removal evidence lacks precheck or backup markers.")
    if removing:
        precheck_index = next(i for i, entry in enumerate(ordered) if entry[0] == "PRECHECKS")
        backup_index = next(i for i, entry in enumerate(ordered) if entry[0] == "BACKUP_PATH")
        first_removal = next(i for i, entry in enumerate(ordered) if entry[0] == "REMOVING")
        if not precheck_index < backup_index < first_removal:
            raise ValueError("Removal precedes prechecks or backup.")
        if any(i >= precheck_index for i, entry in enumerate(ordered) if entry[0] == "ACCOUNT"):
            raise ValueError("Account observations must precede precheck completion.")
    if outcome == "verified":
        reported = _users(one("REMOVED_THIS_RUN"))
        if reason or succeeded_text or not verified or reported != verified or removing != verified:
            raise ValueError("Success result conflicts with per-account verification.")
        if set(accounts) != set(verified) | set(absent):
            raise ValueError("Success does not cover all observed accounts.")
        succeeded = verified
    elif one("REMOVED_THIS_RUN", required=False):
        raise ValueError("Unexpected completed-removal list for this result.")
    if outcome == "already_absent" and (
        not absent or set(accounts) != set(absent) or removing or backup or reason or succeeded_text
    ):
        raise ValueError("No-change result contradicts execution markers.")
    return ExecutionResult(
        vm, outcome, verified, succeeded, tuple(absent), backup, _text(reason), tuple(accounts),
    )


def _headers(sheet) -> dict[str, int]:
    values = [cell.value for cell in sheet[1]]
    if any(not isinstance(value, str) or not value for value in values):
        raise ValueError("Data sheets require nonempty text headers.")
    if len(set(values)) != len(values) or not set(REQUIRED) <= set(values):
        raise ValueError("Expected a local-account workbook with unique required headers.")
    if sheet.tables or sheet.merged_cells or sheet.data_validations.count:
        raise ValueError("Merged cells, tables or validation rules require a reviewed update.")
    if len(sheet.conditional_formatting):
        raise ValueError("Conditional formatting is not supported by this update workflow.")
    return {value: index for index, value in enumerate(values, 1)}


def _set(cell, text: str) -> None:
    cell.value = _text(text)
    cell.data_type = "s"


def _key(sheet, row: int, headers: dict[str, int]) -> tuple[str, str, str]:
    values = tuple(sheet.cell(row, headers[name]).value for name in REQUIRED[:3])
    if any(not isinstance(value, str) or not value.strip() for value in values):
        raise ValueError("Resource keys require Account, Resource Name and Region.")
    return values


def _scope(action: str) -> tuple[str, ...]:
    line = action.splitlines()[0] if action else ""
    if line.startswith("Remove only: "):
        return _users(line.removeprefix("Remove only: "), ";")
    match = re.fullmatch(r"Removed: (.+) \(Run Command verified\)\.", line)
    if match:
        return _users(match[1], ";")
    raise ValueError("Result target lacks an explicit approved local-user action.")


def _done(action: str) -> tuple[str, ...]:
    if action.startswith("Removed: "):
        return _scope(action)
    match = re.fullmatch(
        rf"({USER}) removed \(user-confirmed\)\.", action.splitlines()[0] if action else "",
    )
    return (match[1],) if match else ()


def _read_private(path: str | Path) -> Path:
    source = Path(path)
    if source.is_symlink() or not source.is_file():
        raise ValueError("Input must be a regular, non-symlink file.")
    source = source.resolve()
    if source.is_relative_to(REPOSITORY_ROOT):
        raise ValueError("Operational inputs must be outside the public repository.")
    return source


def _fit(sheet, row: int) -> None:
    lines = 1
    for cell in sheet[row]:
        width = sheet.column_dimensions[cell.column_letter].width or 13
        count = sum(
            max(1, len(wrap(part, max(1, int(width - 2)), break_on_hyphens=False)))
            for part in str(cell.value or "").splitlines()
        )
        lines = max(lines, count)
    size = max(cell.font.sz or 11 for cell in sheet[row])
    height = lines * (size + 3) + 3
    if height > 409:
        raise ValueError(
            "Visible notes exceed Excel row height; shorten them to avoid clipped text."
        )
    sheet.row_dimensions[row].height = height


def _add_columns(sheet, headers) -> dict[str, int]:
    for name in EXTRA:
        if name not in headers:
            column = sheet.max_column + 1
            source = headers["Local User Action"]
            for row in range(1, sheet.max_row + 1):
                new = sheet.cell(row, column)
                new._style = copy(sheet.cell(row, source)._style)
                new.fill = copy(sheet.cell(row, headers["Resource Name"]).fill)
            sheet.cell(1, column)._style = copy(sheet.cell(1, source)._style)
            _set(sheet.cell(1, column), name)
            sheet.column_dimensions[get_column_letter(column)].width = 42
            headers[name] = column
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.print_area = sheet.auto_filter.ref
    return headers


def _load_events(content: bytes | None) -> list[dict]:
    if content is None:
        return []
    try:
        memory = json.loads(content)
        if (
            not isinstance(memory, dict) or memory.get("version") != 1
            or not isinstance(memory.get("events"), list)
        ):
            raise ValueError
        events, ids = [], set()
        for event in memory["events"]:
            if (
                not isinstance(event, dict) or not isinstance(event["id"], str)
                or event["id"] in ids or event["kind"] not in ("result", "blocker")
                or not isinstance(event["key"], list) or len(event["key"]) != 3
                or any(not isinstance(value, str) or not value for value in event["key"])
                or not isinstance(event["received_utc"], str)
                or not isinstance(event["raw_output"], str)
            ):
                raise ValueError
            details = event["details"]
            if (
                not isinstance(details, dict)
                or details["stage"] not in ("Pending review", "Pending remediation")
                or not isinstance(details["scope"], list)
                or not isinstance(details["verified"], list)
                or not isinstance(details["payload"], dict)
                or not isinstance(details["report_state"], dict)
            ):
                raise ValueError
            if details["scope"]:
                _users(",".join(details["scope"]))
            if details["verified"]:
                _users(",".join(details["verified"]))
            if not set(details["verified"]) <= set(details["scope"]):
                raise ValueError
            event["key"] = tuple(event["key"])
            expected = _event_id(event["key"], event["kind"], details["payload"])
            if event["id"] != expected:
                raise ValueError
            ids.add(event["id"])
            events.append(event)
    except (TypeError, KeyError, ValueError):
        raise ValueError("Private execution memory is invalid; restore or reconcile it.") from None
    return events


def _event_id(key, kind, payload) -> str:
    return sha256(json.dumps([key, kind, payload], sort_keys=True).encode("utf-8")).hexdigest()


def _summary(book, rows, states) -> None:
    if "Cover" not in book:
        return
    counts = Counter(row["status"] for row in rows.values())
    completed = sum(
        len(states.get(key, {}).get("verified", _done(row["action"])))
        for key, row in rows.items()
    )
    for cells in book["Cover"]:
        for cell in cells:
            value = cell.value
            if not isinstance(value, str):
                continue
            for cloud in ("AWS", "Azure"):
                if value.startswith(cloud + ": ") and " VMs | " in value:
                    subset = [r for r in rows.values() if r["sheet"].title == cloud + " Data"]
                    stats = Counter(r["status"] for r in subset)
                    _set(cell, f"{cloud}: {len(subset)} VMs | " + " | ".join(
                        f"{status}: {stats[status]}" for status in COLORS
                    ))
                    break
            else:
                if value.startswith("Completed approved local removals:"):
                    _set(cell, "Completed approved local removals: "
                         f"{counts['Remediated']} VMs / {completed} account instances "
                         "(including documented prior completion and verified partial removals). "
                         "This is an execution summary, not closure of all resource risks.")
                elif value.startswith(("Pending remediation:", "Pending remediation / blocked:")):
                    _set(cell, "Pending remediation / blocked: "
                         f"{counts['Pending remediation']} pending remediation; "
                         f"{counts['Blocked']} blocked; {counts['Pending review']} pending review. "
                         "See the pending tabs for blockers and next actions; resource action "
                         "notes retain exclusions.")
                elif value.startswith("Verified only for the "):
                    _set(cell, "Execution evidence is retained separately in private local memory. "
                         "Partial runs and already-absent observations "
                         "are not proof of fully verified removal. See recorded exceptions.")
            if cell.value != value:
                _fit(book["Cover"], cell.row)


def _email(rows, states, ticket: str | None) -> str:
    lines = [
        "Subject: Approved local-account remediation update"
        + (f" | Ticket: {ticket}" if ticket else ""),
        "", "Hi Team,", "",
        "Below is the recorded progress against the approved local-account removal scope. "
        "Imported results are script-reported evidence, not independent live verification.",
        "", "Completed changes",
    ]
    completed = []
    for key, row in rows.items():
        verified = states.get(key, {}).get("verified", _done(row["action"]))
        if verified:
            suffix = " (partial; review required)" if row["status"] != "Remediated" else ""
            completed.append(
                f"- {key[1]} | {key[0]} | {key[2]}: removed {', '.join(verified)}{suffix}."
            )
    lines.extend(completed or ["- No verified account removals are recorded."])
    lines.extend(["", "Blockers and remaining work"])
    pending = []
    for key, row in rows.items():
        if row["status"] not in ("Pending remediation", "Pending review", "Blocked"):
            continue
        scope = states.get(key, {}).get("scope")
        verified = states.get(key, {}).get("verified", [])
        if scope is None:
            try:
                scope = _scope(row["action"])
            except ValueError:
                scope = ()
        remaining = [name for name in scope if name not in verified]
        pending.append(f"- {key[1]} | {key[0]} | {key[2]}: {row['status']}.")
        if remaining:
            pending.append("  Remaining approved users: " + ", ".join(remaining) + ".")
        if row["blocker"]:
            pending.append("  Blocker: " + row["blocker"])
        pending.append("  Next action: " + (row["next"] or (
            "Obtain owner review and explicit scope before making changes."
            if row["status"] == "Pending review"
            else "Confirm the outstanding execution status before scheduling or rerunning changes."
        )))
    lines.extend(pending or ["- No pending or blocked resources are recorded."])
    lines.extend(["", "Scope exceptions and notes"])
    notes = []
    for key, row in rows.items():
        suffix = [
            line for line in row["action"].splitlines()[1:]
            if "Wiz reassessment" not in line and "Wiz assessment" not in line
        ]
        if suffix and row["status"] != "Accepted risk":
            notes.append(f"- {key[1]}: " + " ".join(suffix))
    lines.extend(notes or ["- No additional scope notes are recorded."])
    lines.extend([
        "- Excluded accounts are not counted as removed or as newly accepted risk.",
        "", "Please review the attached report and advise on the blockers and next actions "
        "listed above. Let us know if you need any assistance.",
        "", "Best Regards,", "[Security team]", "",
    ])
    return "\n".join(lines)


def update_report(
    report: str | Path, *, results: Sequence[str | Path] = (),
    blockers: str | Path | None = None, email_output: str | Path,
    account: str | None = None, region: str | None = None, ticket: str | None = None,
    email_only: bool = False,
) -> UpdateResult:
    """Update a supported existing local-account report and write a review-only email.

    No rescan/closure workflow, approval inference, message sending or cloud execution.
    Result names must uniquely match an Azure row (optionally qualified by account/region).
    Close Excel and avoid concurrent writers. A private pre-update backup is always kept.
    """
    if isinstance(results, (str, Path)) or (
        not email_only and not results and blockers is None
    ):
        raise ValueError("Supply result paths and/or a blocker CSV.")
    if email_only and (results or blockers is not None):
        raise ValueError("Email-only mode cannot also change execution or blockers.")
    destination = _output_path(report, datetime.now(UTC).date(), update=True)
    original = destination.read_bytes()
    memory_path = destination.with_suffix(".remediation-memory.json")
    if memory_path.is_symlink():
        raise ValueError("Private memory must not be a symbolic link.")
    old_memory = memory_path.read_bytes() if memory_path.exists() else None
    email = Path(email_output)
    if not email.is_absolute() or email.is_symlink() or email.exists():
        raise ValueError("Email output must be an unused absolute path.")
    email = email.resolve()
    if email.is_relative_to(REPOSITORY_ROOT) or email.suffix.lower() != ".txt":
        raise ValueError("Email output must be a .txt file outside the public repository.")
    if ticket is not None and (not ticket.strip() or "\n" in ticket or "\r" in ticket):
        raise ValueError("Ticket reference must be nonempty single-line text.")
    if ticket:
        _text(ticket)
    try:
        book = load_workbook(BytesIO(original))
    except (BadZipFile, InvalidFileException, ParseError, KeyError, ValueError):
        raise ValueError("Report is not a valid supported XLSX workbook.") from None
    if not all(name in book for name in (*DATA, *PENDING)):
        raise ValueError("Expected AWS/Azure data and both pending tabs; not a credential report.")
    if book._external_links or any(c.data_type == "f" for s in book for row in s for c in row):
        raise ValueError("Formula or external-link workbooks require manual review.")
    headers = {name: _headers(book[name]) for name in (*DATA, *PENDING)}
    if [c.value for c in book[DATA[0]][1]] != [c.value for c in book[DATA[1]][1]]:
        raise ValueError("Consolidated data tabs must have matching headers.")
    rows = {}
    for name in DATA:
        sheet, columns = book[name], headers[name]
        for number in range(2, sheet.max_row + 1):
            key = _key(sheet, number, columns)
            if key in rows:
                raise ValueError("Duplicate consolidated resource key.")
            status = sheet.cell(number, columns["Remediation Status"]).value
            action = sheet.cell(number, columns["Local User Action"]).value or ""
            if status not in COLORS or not isinstance(action, str):
                raise ValueError("Unsupported status or local-user action.")
            rows[key] = {
                "sheet": sheet, "row": number, "status": status, "action": action,
                "blocker": sheet.cell(number, columns["Blocker"]).value or ""
                if "Blocker" in columns else "",
                "next": sheet.cell(number, columns["Next Action"]).value or ""
                if "Next Action" in columns else "",
            }
    events = _load_events(old_memory)
    states = {}
    for event in events:
        if event["key"] not in rows:
            raise ValueError("Execution evidence references a missing resource.")
        states[event["key"]] = event["details"]
    for key, state in states.items():
        if state["scope"] and set(_scope(rows[key]["action"])) != set(state["scope"]):
            raise ValueError(
                "Approved scope changed since the saved execution memory; reconcile it."
            )
        current = {field: rows[key][field] for field in ("status", "action", "blocker", "next")}
        if current != state["report_state"]:
            raise ValueError("Report and private memory disagree; reconcile them before importing.")
    stages = {}
    for name, stage in zip(PENDING, ("Pending review", "Pending remediation"), strict=True):
        for number in range(2, book[name].max_row + 1):
            key = _key(book[name], number, headers[name])
            if key not in rows or key in stages:
                raise ValueError("Pending tabs contain a duplicate or unknown resource.")
            stages[key] = stage
    for key, row in rows.items():
        expected_stage = (
            states.get(key, {}).get("stage", stages.get(key))
            if row["status"] == "Blocked" else row["status"]
        )
        if row["status"] in ("Pending review", "Pending remediation", "Blocked"):
            if expected_stage not in ("Pending review", "Pending remediation") or (
                stages.get(key) != expected_stage
            ):
                raise ValueError("Pending views are inconsistent with consolidated statuses.")
        elif key in stages:
            raise ValueError("Completed or accepted resource appears in a pending view.")
    if email_only:
        _publish(BytesIO(_email(rows, states, ticket).encode("utf-8")), email)
        return UpdateResult(0, 0, None, memory_path)
    operations = []
    for path in results:
        raw_output = _read_private(path).read_text(encoding="utf-8-sig")
        payload = parse_result(raw_output)
        matches = [
            key for key, row in rows.items()
            if row["sheet"].title == "Azure Data" and key[1] == payload.vm
            and (account is None or key[0] == account) and (region is None or key[2] == region)
        ]
        if len(matches) != 1:
            raise ValueError("Result VM is unknown or ambiguous; qualify --account/--region.")
        operations.append((
            matches[0], "result", json.loads(json.dumps(asdict(payload))), raw_output,
        ))
    if blockers is not None:
        with _read_private(blockers).open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            names = ["Account", "Resource Name", "Region", "Status", "Blocker", "Next Action"]
            if reader.fieldnames != names:
                raise ValueError("Blocker CSV headers must be: " + ",".join(names))
            seen = set()
            for item in reader:
                if None in item or any(value is None for value in item.values()):
                    raise ValueError("Malformed blocker CSV row.")
                for value in item.values():
                    _text(value)
                key = tuple(item[name] for name in REQUIRED[:3])
                if key not in rows or key in seen:
                    raise ValueError("Unknown or repeated blocker resource key.")
                if any(op[0] == key for op in operations):
                    raise ValueError("Do not combine result and manual blocker changes for one VM.")
                seen.add(key)
                status = item["Status"]
                if status not in ("Blocked", "Pending review", "Pending remediation"):
                    raise ValueError("Blocker CSV cannot mark remediation complete.")
                if not item["Next Action"].strip() or (
                    status == "Blocked" and not item["Blocker"].strip()
                ) or (status != "Blocked" and item["Blocker"].strip()):
                    raise ValueError("Supply a next action and a blocker only for Blocked status.")
                operations.append((key, "blocker", item, ""))
    if not operations:
        raise ValueError("No execution or blocker records supplied.")
    for name in (*DATA, *PENDING):
        headers[name] = _add_columns(book[name], headers[name])
    event_ids = {event["id"] for event in events}
    imported, duplicates = 0, 0
    received = datetime.now(UTC).isoformat()
    for key, kind, payload, raw_output in operations:
        previous = states.get(key)
        if kind == "blocker":
            previous_payload = previous["payload"] if previous else {}
            if {k: v for k, v in previous_payload.items() if k != "revision"} == payload:
                duplicates += 1
                continue
            payload = {**payload, "revision": len(events)}
        event_id = _event_id(key, kind, payload)
        if event_id in event_ids:
            duplicates += 1
            continue
        row = rows[key]
        stage = previous["stage"] if previous else stages.get(key, "Pending remediation")
        if kind == "result" and (row["status"] == "Accepted risk" or stage == "Pending review"):
            raise ValueError("Execution result targets an unapproved resource.")
        scope = previous["scope"] if previous and previous["scope"] else (
            list(_scope(row["action"])) if kind == "result" else []
        )
        verified = list(previous["verified"]) if previous else list(_done(row["action"]))
        if kind == "result":
            observed = set(payload["accounts"])
            if not observed <= set(scope) or (payload["outcome"] != "stopped"
                                             and observed != set(scope)):
                raise ValueError("Observed accounts do not match the exact approved scope.")
            if row["status"] == "Remediated" and payload["outcome"] != "already_absent":
                raise ValueError("New removal/stop on a completed VM needs manual reconciliation.")
            verified = [name for name in scope if name in set(verified) | set(payload["verified"])]
            suffix = "\n".join(row["action"].splitlines()[1:])
            first = (
                "Removed: " + "; ".join(scope) + " (Run Command verified)."
                if payload["outcome"] == "verified" and set(verified) == set(scope)
                else row["action"].splitlines()[0]
            )
            row["action"] = first + ("\n" + suffix if suffix else "")
            if payload["outcome"] == "stopped":
                row["status"] = "Blocked"
                row["blocker"] = payload["reason"]
                row["next"] = (
                    "Review the stop reason and partial changes with the owner. "
                    "Use read-only verification; do not rerun deletion or assume rollback."
                )
            elif payload["outcome"] == "verified" and set(verified) == set(scope):
                row["status"], row["blocker"], row["next"] = "Remediated", "", ""
            elif row["status"] != "Remediated":
                row["status"] = "Pending remediation"
                row["blocker"] = ""
                row["next"] = (
                    "Targets reported already absent are not proof of removal by this change. "
                    "Reconcile prior execution evidence; do not rerun deletion."
                )
        else:
            if row["status"] in ("Remediated", "Accepted risk"):
                raise ValueError("Manual blockers cannot reopen completed or accepted scope.")
            if row["status"] != "Blocked":
                stage = row["status"]
            if payload["Status"] != "Blocked" and payload["Status"] != stage:
                raise ValueError("Clearing a blocker must retain the original pending stage.")
            row["status"], row["blocker"], row["next"] = (
                payload["Status"], payload["Blocker"], payload["Next Action"],
            )
        details = {
            "stage": stage, "scope": scope, "verified": verified, "payload": payload,
            "report_state": {
                field: row[field] for field in ("status", "action", "blocker", "next")
            },
        }
        states[key] = details
        sheet, number = row["sheet"], row["row"]
        columns = headers[sheet.title]
        for name, value in (
            ("Remediation Status", row["status"]), ("Local User Action", row["action"]),
            ("Blocker", row["blocker"]), ("Next Action", row["next"]),
        ):
            _set(sheet.cell(number, columns[name]), value)
        sheet.cell(number, columns["Remediation Status"]).fill = PatternFill(
            "solid", fgColor=COLORS[row["status"]],
        )
        _fit(sheet, number)
        events.append({
            "id": event_id, "key": key, "kind": kind, "received_utc": received,
            "details": details, "raw_output": raw_output,
        })
        event_ids.add(event_id)
        imported += 1
    for name, stage in zip(PENDING, ("Pending review", "Pending remediation"), strict=True):
        view = book[name]
        view.delete_rows(2, max(0, view.max_row - 1))
        for index in list(view.row_dimensions):
            if index > 1:
                del view.row_dimensions[index]
        columns = headers[name]
        for key, row in rows.items():
            actual_stage = (
                states.get(key, {}).get("stage", stages.get(key))
                if row["status"] == "Blocked" else row["status"]
            )
            if actual_stage != stage or row["status"] == "Remediated":
                continue
            number = view.max_row + 1
            for title, column in columns.items():
                source_column = headers[row["sheet"].title].get(title)
                if source_column is None:
                    raise ValueError("Pending view has a header missing from consolidated data.")
                source = row["sheet"].cell(row["row"], source_column)
                cell = view.cell(number, column)
                cell.value, cell.data_type = source.value, source.data_type
                cell._style = copy(source._style)
                cell.comment = copy(source.comment)
                if source.hyperlink:
                    cell._hyperlink = copy(source.hyperlink)
                    cell._hyperlink.ref = cell.coordinate
                cell.fill = PatternFill("solid", fgColor=(
                    COLORS[row["status"]] if title == "Remediation Status"
                    else "FFD9E8F5" if number % 2 == 0 else "FFFFFFFF"
                ))
            _fit(view, number)
        view.auto_filter.ref = f"A1:{get_column_letter(view.max_column)}{view.max_row}"
        view.print_area = view.auto_filter.ref
    _summary(book, rows, states)
    email_text = _email(rows, states, ticket)
    stream = BytesIO()
    book.save(stream)
    check = load_workbook(BytesIO(stream.getvalue()))
    if check._external_links or any(c.data_type == "f" for s in check for row in s for c in row):
        raise ValueError("Unexpected formula or external workbook link after update.")
    memory_bytes = json.dumps({"version": 1, "events": events}, indent=2).encode("utf-8")
    _load_events(memory_bytes)
    backup = destination.with_name(f"{destination.stem}.before-{uuid4().hex}.xlsx")
    _publish(BytesIO(original), backup)
    _publish(BytesIO(email_text.encode("utf-8")), email)
    memory_saved = False
    try:
        _publish(
            BytesIO(memory_bytes), memory_path,
            expected_digest=sha256(old_memory).digest() if old_memory is not None else None,
        )
        memory_saved = True
        _publish(stream, destination, expected_digest=sha256(original).digest())
    except (OSError, ValueError):
        if memory_saved:
            if old_memory is None:
                if memory_path.read_bytes() != memory_bytes:
                    raise ValueError(
                        "Concurrent memory change; reconcile the backup before retry."
                    ) from None
                memory_path.unlink()
            else:
                _publish(
                    BytesIO(old_memory), memory_path,
                    expected_digest=sha256(memory_bytes).digest(),
                )
        if email.is_file() and email.read_bytes() == email_text.encode("utf-8"):
            email.unlink()
        raise
    return UpdateResult(imported, duplicates, backup, memory_path)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--result", action="append", type=Path, default=[])
    parser.add_argument("--blockers", type=Path)
    parser.add_argument("--email-output", type=Path, required=True)
    parser.add_argument("--account", help="Exact report account to disambiguate result VM names")
    parser.add_argument("--region", help="Exact report region to disambiguate result VM names")
    parser.add_argument("--ticket", help="Existing ticket reference; omitted when not supplied")
    parser.add_argument(
        "--email-only", action="store_true",
        help="Draft from current records without changing Excel",
    )
    args = parser.parse_args(argv)
    try:
        result = update_report(
            args.report, results=args.result, blockers=args.blockers,
            email_output=args.email_output, account=args.account, region=args.region,
            ticket=args.ticket, email_only=args.email_only,
        )
    except (OSError, ValueError, csv.Error) as error:
        print(f"Update failed: {error}", file=sys.stderr)
        return 1
    if args.email_only:
        print("Draft generated from current records; report unchanged.")
    else:
        print(
            f"Report updated. Imported: {result.imported}; duplicates skipped: {result.duplicates}."
        )
        print(f"Private backup: {result.backup_path}")
        print(f"Private execution memory: {result.memory_path}")
    print("Stakeholder draft saved; review before sending. No cloud actions or messages sent.")
    return 0
