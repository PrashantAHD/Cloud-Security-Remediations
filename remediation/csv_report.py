"""Read private Wiz graph CSV exports directly into a text-safe evidence workbook.

No cases, database, network calls, or Wiz lifecycle transitions are involved.
``generate_report`` is the programmatic entry point; ``python -m`` exposes the CLI.
"""

import argparse
import csv
import json
import os
import re
import sys
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile

from defusedxml.ElementTree import ParseError
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.exceptions import InvalidFileException

from remediation.csv_layout import cover, remediation_page, resource_table
from remediation.issue_metadata import enrich_records
from remediation.workflow import Tracker, load_tracker

PRINCIPAL = "(SERVICE_ACCOUNT | USER_ACCOUNT)."
KEY = "ACCESS_KEY."
PERMISSION = "ACCESS_ROLE_PERMISSION."
SCOPE = "(CLOUD_ORGANIZATION | KUBERNETES_CLUSTER | NAMESPACE | SUBSCRIPTION)."
REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SEVERITIES = ("Critical", "High", "Medium", "Low", "Informational")
NOT_SUPPLIED = "Not supplied"
UNKNOWN = "Unknown"
DEFAULT_TITLE = "Wiz identity credential evidence"
SHEETS = ["Cover", "AWS Data", "Azure Data", "Remediation"]
CREATOR = "Cloud Security Remediations"
_ROW_LIMIT = 1048575
_COLUMN_LIMIT = 16384
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}\Z")
_TIMESTAMP = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})\Z")


@dataclass(frozen=True)
class ReportResult:
    """Saved path and counts; IDs/names are deliberately absent."""

    path: Path
    source_rows: int
    unique_principals: int
    unique_credentials: int
    grouped_scopes: int
    aws_principals: int
    aws_credentials: int
    azure_principals: int
    azure_credentials: int
    unique_issues: int | None = None


@dataclass
class _Evidence:
    records: list[dict[str, str]]
    source: list[dict[str, str]]
    source_fields: list[tuple[str, str]]


def _validate_text(value: str) -> str:
    if ILLEGAL_CHARACTERS_RE.search(value) or any(
        0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in value
    ):
        raise ValueError("Input contains characters that cannot be represented in XLSX.")
    return value


def _as_of(value: str | date | None) -> date:
    if value is None:
        return datetime.now(UTC).date()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, str) and _DATE.fullmatch(value):
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    raise ValueError("As-of date must be a valid YYYY-MM-DD date.")


def _utc_date(value: str, field: str) -> date | None:
    if not value.strip():
        return None
    try:
        if _DATE.fullmatch(value):
            return date.fromisoformat(value)
        if _TIMESTAMP.fullmatch(value):
            return datetime.fromisoformat(value).astimezone(UTC).date()
    except ValueError:
        pass
    raise ValueError(f"{field} must be an ISO date or a timezone-qualified ISO timestamp.")


def _cloud(row: dict[str, str]) -> str:
    values = [
        row.get(SCOPE + "cloudPlatform", "").strip(),
        row.get(PERMISSION + "cloudPlatform", "").strip(),
    ]
    normalized = []
    for value in values:
        if value:
            cloud = {"aws": "AWS", "azure": "Azure"}.get(value.casefold())
            if cloud is None:
                raise ValueError("Unsupported cloud; only AWS and Azure exports are supported.")
            normalized.append(cloud)
    if not normalized:
        raise ValueError("Cloud must be supplied by scope or permission cloudPlatform.")
    if len(set(normalized)) != 1:
        raise ValueError("Scope and permission cloudPlatform conflict.")
    return normalized[0]


def _scope_type(row: dict[str, str]) -> str:
    # Keep the actual type evidence: a subscription-shaped graph node can be a tenant.
    return "; ".join(
        f"{label}: {row.get(SCOPE + field, '').strip() or UNKNOWN}"
        for field, label in (
            ("RowObjectType", "RowObjectType"),
            ("nativeType", "Native type"),
            ("subscriptionId", "Subscription ID"),
        )
    )


def _access_types(value: str) -> list[str]:
    if not value.strip():
        return []
    if value.lstrip().startswith("["):
        try:
            items = json.loads(value)
        except ValueError:
            raise ValueError("Permission accessTypes contains an invalid JSON array.") from None
        if not isinstance(items, list) or any(not isinstance(item, str) for item in items):
            raise ValueError("Permission accessTypes must contain text values.")
        return [_validate_text(item) for item in items if item.strip()]
    return [value]


def _record(row: dict[str, str], as_of: date) -> dict[str, str]:
    rotated = row.get(KEY + "rotatedAt", "")
    expiry = row.get(KEY + "validBefore", "")
    rotated_date = _utc_date(rotated, "rotatedAt")
    expiry_date = _utc_date(expiry, "validBefore")
    if rotated_date is not None and rotated_date > as_of:
        raise ValueError("rotatedAt is after the as-of UTC date.")
    return {
        "cloud": _cloud(row),
        "scope_id": row[SCOPE + "ID"],
        "scope_name": row.get(SCOPE + "Name", "") or NOT_SUPPLIED,
        "scope_external_id": row.get(SCOPE + "externalId", "") or NOT_SUPPLIED,
        "scope_type": _scope_type(row),
        "principal": row.get(PRINCIPAL + "Name", "") or NOT_SUPPLIED,
        "principal_native_id": row.get(PRINCIPAL + "providerUniqueId", "") or NOT_SUPPLIED,
        "principal_id": row[PRINCIPAL + "ID"],
        "principal_type": row.get(PRINCIPAL + "RowObjectType", "") or UNKNOWN,
        "credential_name": row.get(KEY + "Name", "") or NOT_SUPPLIED,
        "credential_id": row[KEY + "ID"],
        "key_active": row.get(KEY + "active", "").strip() or UNKNOWN,
        "high_privileges": row.get(PRINCIPAL + "hasHighPrivileges", "").strip() or UNKNOWN,
        "inactive": row.get(PRINCIPAL + "inactiveInLast90Days", "").strip() or UNKNOWN,
        "rotated_at": rotated if rotated.strip() else UNKNOWN,
        "rotated_date": rotated_date.isoformat() if rotated_date else UNKNOWN,
        "age_days": str((as_of - rotated_date).days) if rotated_date else UNKNOWN,
        "valid_before": expiry if expiry.strip() else NOT_SUPPLIED,
        "expiry_date": expiry_date.isoformat() if expiry_date else NOT_SUPPLIED,
        "days_until_expiry": str((expiry_date - as_of).days) if expiry_date else NOT_SUPPLIED,
        "status": "Draft",
    }


def _read_exports(paths: Sequence[Path], as_of: date) -> _Evidence:
    groups: dict[tuple[str, str, str, str], dict[str, str]] = {}
    group_facts: dict[tuple[str, str, str, str], dict[str, str]] = {}
    first_rows: dict[tuple[str, str, str, str], str] = {}
    entity_facts: dict[tuple[str, str, str], tuple[dict[str, str], str]] = {}
    permissions: dict[tuple[str, str, str, str], dict[str, list[str]]] = {}
    source: list[dict[str, str]] = []
    headers: dict[str, str] = {}
    source_fields = [
        ("source_file", "Source File"),
        ("source_input", "Source Input"),
        ("source_row", "Source Row"),
    ]
    required = (PRINCIPAL + "ID", KEY + "ID", SCOPE + "ID")
    for index, path in enumerate(paths, 1):
        _validate_text(path.name)
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, strict=True)
            try:
                fieldnames = next(reader)
            except StopIteration:
                raise ValueError("CSV export is empty.") from None
            if not fieldnames or any(not field.strip() for field in fieldnames):
                raise ValueError("CSV headers must not be empty.")
            if len(set(fieldnames)) != len(fieldnames):
                raise ValueError("CSV headers must be unique.")
            if any(field not in fieldnames for field in required):
                raise ValueError(
                    "Unsupported CSV schema: required graph grouping ID headers missing."
                )
            if not any(
                prefix + "cloudPlatform" in fieldnames for prefix in (SCOPE, PERMISSION)
            ):
                raise ValueError("Unsupported CSV schema: cloudPlatform header missing.")
            for field in fieldnames:
                _validate_text(field)
                if len(field) > 32767:
                    raise ValueError("CSV header exceeds Excel's text limit.")
                if field not in headers:
                    key = f"original_{len(headers)}"
                    headers[field] = key
                    source_fields.append((key, field))
            if len(source_fields) > _COLUMN_LIMIT:
                raise ValueError("Source evidence exceeds Excel's column limit.")
            count = 0
            while True:
                line = reader.line_num + 1
                try:
                    values = next(reader)
                except StopIteration:
                    break
                if len(values) != len(fieldnames):
                    raise ValueError("CSV contains a ragged row.")
                row = dict(zip(fieldnames, map(_validate_text, values), strict=True))
                if any(not row[field].strip() for field in required):
                    raise ValueError("CSV grouping IDs must not be blank.")
                count += 1
                if len(source) >= _ROW_LIMIT:
                    raise ValueError("Source evidence exceeds Excel's row limit.")
                source.append({
                    "source_file": path.name,
                    "source_input": str(index),
                    "source_row": str(line),
                    **{headers[field]: value for field, value in row.items()},
                })
                record = _record(row, as_of)
                group = tuple(record[field] for field in (
                    "cloud", "principal_id", "credential_id", "scope_id"
                ))
                facts = {
                    field: value.strip() for field, value in row.items()
                    if field.startswith((PRINCIPAL, KEY, SCOPE)) and value.strip()
                }
                reference = f"input {index}, row {line}"
                for prefix, identity in (
                    (PRINCIPAL, record["principal_id"]),
                    (KEY, record["credential_id"]),
                    (SCOPE, record["scope_id"]),
                ):
                    entity = (record["cloud"], prefix, identity)
                    observed = {
                        field: value for field, value in facts.items() if field.startswith(prefix)
                    }
                    if prefix == KEY:
                        observed["owner_principal_id"] = record["principal_id"]
                    if entity in entity_facts and entity_facts[entity][0] != observed:
                        raise ValueError(
                            "Conflicting entity evidence between "
                            f"{entity_facts[entity][1]} and {reference}; "
                            "review snapshot drift before combining exports."
                        )
                    entity_facts.setdefault(entity, (observed, reference))
                if group in groups and (
                    groups[group] != record or group_facts[group] != facts
                ):
                    raise ValueError(
                        "Conflicting credential-scope evidence between "
                        f"{first_rows[group]} and {reference}; "
                        "review snapshot drift before combining exports."
                    )
                groups[group] = record
                group_facts[group] = facts
                first_rows.setdefault(group, reference)
                merged = permissions.setdefault(group, {
                    "permission_names": [], "permission_access_types": [], "source_rows": [],
                })
                for key, values in (
                    ("permission_names", [row.get(PERMISSION + "Name", "")]),
                    ("permission_access_types", _access_types(
                        row.get(PERMISSION + "accessTypes", "")
                    )),
                ):
                    for value in values:
                        if value.strip() and value not in merged[key]:
                            merged[key].append(value)
                merged["source_rows"].append(f"Input {index}: {path.name}, row {line}")
            if not count:
                raise ValueError("CSV export contains headers but no evidence rows.")
    for group, record in groups.items():
        record.update({
            field: "\n".join(values) or NOT_SUPPLIED
            for field, values in permissions[group].items()
        })
    return _Evidence(list(groups.values()), source, source_fields)


def _counts(
    evidence: _Evidence, path: Path, unique_issues: int | None = None
) -> ReportResult:
    records = evidence.records
    principals = {(row["cloud"], row["principal_id"]) for row in records}
    credentials = {(row["cloud"], row["credential_id"]) for row in records}
    return ReportResult(
        path=path,
        source_rows=len(evidence.source),
        unique_principals=len(principals),
        unique_credentials=len(credentials),
        grouped_scopes=len(records),
        aws_principals=sum(cloud == "AWS" for cloud, _ in principals),
        aws_credentials=sum(cloud == "AWS" for cloud, _ in credentials),
        azure_principals=sum(cloud == "Azure" for cloud, _ in principals),
        azure_credentials=sum(cloud == "Azure" for cloud, _ in credentials),
        unique_issues=unique_issues,
    )


def _workbook(
    evidence: _Evidence,
    result: ReportResult,
    as_of: date,
    title: str,
    severity: str,
    rule_id: str,
    notes: str,
    issues: list[dict[str, str]] | None = None,
    workflow: Tracker | None = None,
) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = CREATOR
    workbook.properties.title = title
    cover(
        workbook, evidence.records, title=title, as_of=as_of, source_rows=result.source_rows,
        severity=severity, rule_id=rule_id, issues=issues,
        workflow_text=workflow.cover_text() if workflow is not None else "",
        workflow_context=workflow.display() if workflow is not None else "",
    )
    for name, subset in (
        ("AWS Data", [row for row in evidence.records if row["cloud"] == "AWS"]),
        ("Azure Data", [row for row in evidence.records if row["cloud"] == "Azure"]),
    ):
        resource_table(workbook, name, subset)
    remediation_page(workbook, notes=notes)
    workbook.active = 0
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    check = load_workbook(stream, read_only=True, keep_links=False)
    try:
        if check.sheetnames != SHEETS:
            raise ValueError("Generated workbook does not have the required four tabs.")
        for sheet in check:
            for row in sheet:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Generated workbook contains an unexpected formula.")
    finally:
        check.close()
    stream.seek(0)
    return stream


def _output_path(output: str | Path | None, as_of: date, *, update: bool = False) -> Path:
    if output is None:
        if update:
            raise ValueError("Update requires an explicit --output path.")
        local = os.environ.get("LOCALAPPDATA")
        if not local:
            raise ValueError("LOCALAPPDATA is unavailable; supply an absolute --output path.")
        output = (
            Path(local) / "CloudSecurityRemediations" / "reports"
            / f"credential-evidence-{as_of.isoformat()}-{uuid4().hex}.xlsx"
        )
    path = Path(output)
    if not path.is_absolute():
        raise ValueError("Output path must be absolute.")
    if path.is_symlink():
        raise ValueError("Output must not be a symbolic link.")
    if path.exists() and not update:
        raise ValueError("Output already exists; refusing to overwrite.")
    path = path.resolve()
    if path.is_relative_to(REPOSITORY_ROOT):
        raise ValueError("Output must be outside the repository.")
    if path.suffix.casefold() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension.")
    if update and not path.is_file():
        raise ValueError("Update requires an existing regular XLSX file.")
    if path.exists() and not update:
        raise ValueError("Output already exists; refusing to overwrite.")
    return path


def _update_digest(output: Path) -> bytes:
    content = output.read_bytes()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, ParseError, KeyError, ValueError):
        raise ValueError("Update target is not a valid tool-generated workbook.") from None
    try:
        if workbook.properties.creator != CREATOR or workbook.sheetnames != SHEETS:
            raise ValueError("Update requires an existing four-tab workbook from this tool.")
    finally:
        workbook.close()
    return sha256(content).digest()


def _publish(stream: BytesIO, output: Path, *, expected_digest: bytes | None = None) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    staging = output.with_name(f".{uuid4().hex}.partial")
    try:
        with staging.open("xb") as handle:
            handle.write(stream.getbuffer())
            handle.flush()
            os.fsync(handle.fileno())
        if expected_digest is None:
            # Exclusive creation cannot replace a concurrent writer's file.
            os.link(staging, output)
        else:
            if output.is_symlink() or sha256(output.read_bytes()).digest() != expected_digest:
                raise ValueError("Update target changed during generation; no update was made.")
            os.replace(staging, output)
    finally:
        staging.unlink(missing_ok=True)


def _issue_context(
    issues: list[dict[str, str]], title: str | None, severity: str, rule_id: str
) -> tuple[str, str, str]:
    titles = {
        issue["issue_title"] for issue in issues
        if issue["issue_title"].strip() not in ("", NOT_SUPPLIED, UNKNOWN)
    }
    if len(titles) > 1:
        raise ValueError("Issue export contains multiple finding titles; use one finding.")
    source_title = next(iter(titles), DEFAULT_TITLE)
    if title is not None and titles and title != source_title:
        raise ValueError("Provided title conflicts with the issue export.")
    source_rule = issues[0]["issue_control_id"]
    if rule_id != NOT_SUPPLIED and rule_id != source_rule:
        raise ValueError("Provided rule ID conflicts with the issue export.")
    severities = sorted({issue["issue_severity"] for issue in issues})
    if severity != NOT_SUPPLIED and severities != [severity]:
        raise ValueError("Provided severity conflicts with or is unverified by the issue export.")
    source_severity = severities[0] if len(severities) == 1 else "Mixed: " + ", ".join(severities)
    return title or source_title, source_severity, source_rule


def generate_report(
    paths: Sequence[str | Path],
    *,
    output: str | Path | None = None,
    as_of: str | date | None = None,
    title: str | None = None,
    severity: str = NOT_SUPPLIED,
    rule_id: str = NOT_SUPPLIED,
    notes: str | Path | None = None,
    issues: Sequence[str | Path] | None = None,
    update: bool = False,
    workflow: str | Path | None = None,
) -> ReportResult:
    """Publish four tabs outside this repository and return non-identifying counts.

    ``paths`` contains one or more UTF-8/BOM CSV paths. ``notes`` is a local UTF-8
    text-file path, not inline text. ``as_of`` accepts a date or YYYY-MM-DD string.
    Invalid/conflicting evidence raises ValueError; file/CSV errors may also
    propagate. No output is published until the entire workbook is validated.
    Optional issue exports must describe the same finding and match all graph
    principals by exact IDs. Explicit updates regenerate, not merge, the report:
    manual workbook edits are replaced. Close Excel and avoid concurrent writers.
    ``workflow`` optionally snapshots a private tracker with the same control ID;
    scope still requires human review. Report generation never advances it.
    """
    if not paths or isinstance(paths, (str, Path)):
        raise ValueError("Supply one or more CSV paths as a sequence.")
    if severity not in (*SEVERITIES, NOT_SUPPLIED):
        raise ValueError("Severity must be Critical, High, Medium, Low or Informational.")
    day = _as_of(as_of)
    if isinstance(issues, (str, Path)):
        raise ValueError("Supply issue CSV paths as a sequence.")
    destination = _output_path(output, day, update=update)
    expected_digest = _update_digest(destination) if update else None
    for value in (title or DEFAULT_TITLE, rule_id):
        _validate_text(value)
    note_text = (
        _validate_text(Path(notes).read_text(encoding="utf-8-sig")) if notes else ""
    )
    previous_limit = csv.field_size_limit()
    issue_records = None
    try:
        csv.field_size_limit(2**31 - 1)
        evidence = _read_exports([Path(path) for path in paths], day)
        if issues is not None:
            if not issues:
                raise ValueError("Supply one or more issue CSV paths.")
            enriched = enrich_records(evidence.records, issues)
            evidence.records = enriched.records
            issue_records = enriched.issues
            title, severity, rule_id = _issue_context(issue_records, title, severity, rule_id)
    finally:
        csv.field_size_limit(previous_limit)
    tracker = load_tracker(workflow) if workflow is not None else None
    if tracker is not None and (
        rule_id == NOT_SUPPLIED or tracker.control_id != rule_id
    ):
        raise ValueError(
            "Workflow control ID must match the report's explicit or verified rule ID."
        )
    result = _counts(
        evidence, destination, len(issue_records) if issue_records is not None else None
    )
    stream = _workbook(
        evidence, result, day, title or DEFAULT_TITLE, severity, rule_id, note_text, issue_records,
        tracker,
    )
    _publish(stream, destination, expected_digest=expected_digest)
    return result


def main(argv: Sequence[str] | None = None) -> int:
    """Run the local CSV-only CLI without printing source identifiers or names."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", help="Local UTF-8 Wiz graph CSV exports")
    parser.add_argument("--output", help="Absolute .xlsx path outside the repository")
    parser.add_argument(
        "--update", action="store_true",
        help="Regenerate an existing tool workbook at --output; replaces manual workbook edits",
    )
    parser.add_argument(
        "--issues", action="append",
        help="Issue-level CSV for the same finding; repeat for multiple files",
    )
    parser.add_argument("--as-of", help="UTC date YYYY-MM-DD (default: UTC today)")
    parser.add_argument("--title", help="Finding title; otherwise taken from the issue export")
    parser.add_argument("--severity", choices=SEVERITIES, default=NOT_SUPPLIED)
    parser.add_argument("--rule-id", default=NOT_SUPPLIED)
    parser.add_argument("--notes", help="Local UTF-8 text file kept in a Remediation title comment")
    parser.add_argument("--workflow", help="Private tracker JSON; snapshot its stage on the cover")
    args = parser.parse_args(argv)
    if args.update:
        print(
            "Update regenerates all four tabs; manual workbook edits will be replaced. "
            "Close Excel and avoid concurrent writers.",
            file=sys.stderr,
        )
    try:
        result = generate_report(
            args.paths, output=args.output, as_of=args.as_of, title=args.title,
            severity=args.severity, rule_id=args.rule_id, notes=args.notes,
            issues=args.issues, update=args.update, workflow=args.workflow,
        )
    except ValueError as error:
        print(f"Report not created: {error}", file=sys.stderr)
        return 1
    except (OSError, UnicodeError, csv.Error, OverflowError):
        print("Report not created: could not read valid UTF-8 CSV/notes or publish output.",
              file=sys.stderr)
        return 1
    print(f"Saved: {result.path}")
    print(
        f"Source rows: {result.source_rows}; unique principals: {result.unique_principals}; "
        f"unique credentials: {result.unique_credentials}; grouped scopes: {result.grouped_scopes}"
    )
    print(
        f"AWS principals: {result.aws_principals}; AWS credentials: {result.aws_credentials}; "
        f"Azure principals: {result.azure_principals}; "
        f"Azure credentials: {result.azure_credentials}"
    )
    issue_count = result.unique_issues if result.unique_issues is not None else UNKNOWN
    print(f"Unique Wiz issues: {issue_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
