"""Reusable, read-only credential report layouts; source CSVs remain the evidence."""

from collections import Counter
from collections.abc import Mapping, Sequence
from datetime import date
from textwrap import wrap
from unicodedata import east_asian_width

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.views import Selection
from openpyxl.worksheet.worksheet import Worksheet

from remediation.issue_metadata import validate_wiz_url

NAVY = "1A3D5C"
BLUE = "003D7A"
LIGHT_BLUE = "E8EFF7"
STRIPE = "D9E8F5"
RED = "CC3333"
ORANGE = "FF9900"
GREEN = "00A651"
MAX_TEXT = 32767
MAX_HEIGHT = 409
MAX_ROWS = 1048576
THIN_BORDER = Border(
    left=Side(style="thin", color="D2DCE5"),
    right=Side(style="thin", color="D2DCE5"),
    top=Side(style="thin", color="D2DCE5"),
    bottom=Side(style="thin", color="D2DCE5"),
)
TABLE_FIELDS = (
    ("account", "Account / Scope", 28),
    ("principal", "Resource Name", 35),
    ("credential", "Credential", 35),
    ("rotated", "Last Rotated", 15),
    ("age", "Age (Days)", 12),
    ("expiry", "Expires", 15),
    ("permissions", "Permission Summary", 45),
    ("inactivity", "Inactivity Status", 18),
    ("status", "Remediation Status", 18),
)
COMMENT_FIELDS = (
    ("cloud", "Cloud"),
    ("scope_name", "Scope name"),
    ("scope_external_id", "Scope native ID"),
    ("scope_id", "Scope Wiz ID"),
    ("scope_type", "Scope type"),
    ("principal_id", "Resource Wiz ID"),
    ("principal_native_id", "Resource native ID"),
    ("principal_type", "Resource type"),
    ("credential_id", "Credential ID"),
    ("key_active", "Credential active (source)"),
    ("high_privileges", "Highly privileged (source)"),
    ("inactive", "Inactivity flag (source)"),
    ("rotated_at", "Last rotated (source)"),
    ("valid_before", "Expiry (source)"),
    ("days_until_expiry", "Days until expiry (as of report)"),
    ("permission_access_types", "Permission access types"),
    ("source_rows", "Source CSV rows"),
    ("issue_id", "Issue ID"),
    ("issue_control_id", "Issue control ID"),
    ("issue_title", "Issue title"),
    ("issue_severity", "Issue severity"),
    ("issue_status", "Wiz lifecycle status"),
    ("issue_status_raw", "Wiz lifecycle status (raw)"),
    ("issue_resource_type", "Issue resource type"),
    ("issue_resource_name", "Issue resource name"),
    ("issue_native_id", "Issue resource native ID"),
    ("issue_subscription_id", "Issue subscription ID"),
    ("issue_subscription_name", "Issue subscription name"),
    ("issue_created_at", "Issue created at"),
    ("issue_updated_at", "Issue updated at (not a fresh credential observation)"),
    ("issue_status_changed_at", "Issue status changed at"),
    ("issue_reopened_at", "Issue reopened at"),
    ("issue_assignee", "Issue assignee"),
    ("issue_due_at", "Issue due at (source; not an invented SLA)"),
    ("issue_ticket_references", "Issue ticket references"),
    ("issue_url", "Source issue URL"),
    ("issue_recommendation", "Source issue recommendation (not approved guidance)"),
    ("issue_description", "Source issue description"),
    ("issue_scope_note", "Issue and graph scope context"),
)


def _value(record: Mapping[str, str], key: str, default: str = "Unknown") -> str:
    value = record.get(key)
    return str(value) if value is not None and str(value).strip() else default


def _check_text(text: str, *, notes: bool = False) -> None:
    if len(text) > MAX_TEXT:
        label = "Review notes" if notes else "Report text"
        raise ValueError(f"{label} exceeds the Excel 32767-character limit.")
    if ILLEGAL_CHARACTERS_RE.search(text) or any(
        0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in text
    ):
        raise ValueError("Report text contains characters not supported by XLSX.")


def follow_up_columns(approval_reference: str | None = None) -> tuple[str, ...]:
    """Expose follow-up columns only for an operator-reviewed approval response."""
    if approval_reference is None:
        return ()
    if not isinstance(approval_reference, str) or not approval_reference.strip():
        raise ValueError("A nonblank reviewed approval-response reference is required.")
    _check_text(approval_reference)
    return ("Blocker", "Next Action")


def _literal(sheet: Worksheet, row: int, column: int, text: str) -> None:
    _check_text(text)
    cell = sheet.cell(row, column)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"


def _comment(text: str) -> Comment:
    _check_text(text)
    return Comment(text, "Report context")


def _height(text: str, width: float, size: int = 10, minimum: float = 24) -> float:
    # Excel does not auto-fit wrapped content reliably; reserve font-scaled line space.
    _check_text(text)
    chars = max(1, int((width - 3) * 8 / size))
    measured = "".join(
        char * 2 if char in "MW@%&" or east_asian_width(char) in ("W", "F") else char
        for char in text
    )
    lines = sum(
        max(1, len(wrap(line.expandtabs(4), width=chars, break_on_hyphens=False)))
        for line in measured.split("\n")
    )
    height = max(minimum, lines * (size + 4) + 10)
    if height > MAX_HEIGHT:
        raise ValueError("Report text exceeds the Excel 409-point row height limit.")
    return height


def _print_setup(sheet: Worksheet, *, landscape: bool = False) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if landscape else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.5
    sheet.page_margins.header = sheet.page_margins.footer = 0.15
    sheet.oddFooter.left.text = "Confidential - internal review"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.print_options.horizontalCentered = True


def _page(workbook: Workbook, name: str) -> Worksheet:
    sheet = workbook.create_sheet(name)
    sheet.column_dimensions["A"].width = 100
    _print_setup(sheet)
    return sheet


def _block(
    sheet: Worksheet,
    row: int,
    text: str,
    *,
    fill: str = "FFFFFF",
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
    center: bool = False,
    border: bool = False,
    minimum: float = 24,
) -> int:
    height = _height(text, 100, size, minimum)
    _literal(sheet, row, 1, text)
    cell = sheet.cell(row, 1)
    cell.font = Font(name="Calibri", size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center" if center or bold else "top",
        wrap_text=True,
    )
    if border:
        cell.border = THIN_BORDER
    sheet.row_dimensions[row].height = height
    return row + 1


def _section(sheet: Worksheet, row: int, text: str, fill: str = GREEN) -> int:
    return _block(sheet, row, text, fill=fill, size=12, bold=True, color="FFFFFF")


def _metric(sheet: Worksheet, row: int, text: str) -> None:
    _block(sheet, row, text, fill=LIGHT_BLUE, size=11, bold=True, color=BLUE, border=True)


def _title(sheet: Worksheet, title: str, subtitle: str) -> None:
    _block(sheet, 1, title, fill=NAVY, size=28,
           bold=True, color="FFFFFF", center=True, minimum=70.05)
    _block(sheet, 3, subtitle, fill="2E5266", size=16,
           bold=True, color="FFFFFF", center=True, minimum=30)


def _number(value: str) -> int | None:
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def _required_id(record: Mapping[str, str], key: str) -> str:
    value = _value(record, key, "")
    if not value:
        raise ValueError(f"Report records require a nonblank {key} identifier.")
    return value


def _identity(record: Mapping[str, str]) -> tuple[str, str]:
    return _required_id(record, "cloud"), _required_id(record, "principal_id")


def _credential(record: Mapping[str, str]) -> tuple[str, str]:
    return _required_id(record, "cloud"), _required_id(record, "credential_id")


def _severity(value: str) -> str:
    normalized = value.strip().casefold()
    return {
        "critical": "Critical", "high": "High", "medium": "Medium", "low": "Low",
        "informational": "Informational", "info": "Informational",
    }.get(normalized, "Unknown")


def _issues(
    records: Sequence[Mapping[str, str]], issues: Sequence[Mapping[str, str]] | None,
) -> dict[str, Mapping[str, str]]:
    result = {}
    for issue in issues if issues is not None else records:
        issue_id = _value(issue, "issue_id", "")
        if issue_id:
            result[issue_id] = issue
    return result


def _counts(values: Sequence[str]) -> str:
    return ", ".join(f"{key}: {value}" for key, value in sorted(Counter(values).items()))


def cover(
    workbook: Workbook,
    records: Sequence[Mapping[str, str]],
    *,
    title: str,
    as_of: date,
    source_rows: int,
    severity: str,
    rule_id: str,
    issues: Sequence[Mapping[str, str]] | None = None,
    workflow_text: str = "",
    workflow_context: str = "",
) -> None:
    sheet = _page(workbook, "Cover")
    _title(sheet, "Privileged Credentials Report", "AWS & Azure Multi-Cloud Assessment")
    _block(sheet, 5, "Executive Summary", fill=BLUE, size=13, bold=True, color="FFFFFF")
    _metric(sheet, 6, title)
    for row, cloud in ((7, "AWS"), (8, "Azure")):
        subset = [r for r in records if _value(r, "cloud").casefold() == cloud.casefold()]
        _metric(sheet, row, f"{cloud}: {len({_identity(r) for r in subset})} identities | "
                f"{len({_credential(r) for r in subset})} credentials")
    unique_issues = _issues(records, issues)
    issue_count = (
        f"{len(unique_issues)} Wiz issue{'s' if len(unique_issues) != 1 else ''}"
    ) if unique_issues else (
        "0 Wiz issues" if issues is not None else "Wiz issue count unknown"
    )
    _metric(sheet, 9,
            f"Total: {len({_identity(r) for r in records})} identities | "
            f"{len({_credential(r) for r in records})} credentials | {issue_count}")
    sheet["A9"].comment = _comment(
        f"{source_rows} source CSV rows; {len(records)} credential-scope rows. "
        "These are evidence counts, not additional issues."
    )
    _metric(sheet, 10, f"Report date: {as_of:%Y-%m-%d} (UTC) | Source: Wiz export")
    sheet["A10"].comment = _comment(
        "The report date is not evidence freshness. Source observations are not live checks; "
        "current configuration still needs confirmation."
    )
    unique_keys = {_credential(r): r for r in records}
    aged = sum(
        _value(r, "key_active").casefold() == "true"
        and _value(r, "high_privileges").casefold() == "true"
        and (age := _number(_value(r, "age_days"))) is not None and age > 365
        for r in unique_keys.values()
    )
    expiring = [
        r for r in unique_keys.values()
        if (days := _number(_value(r, "days_until_expiry"))) is not None and 0 <= days <= 60
    ]
    expired = sum(
        (days := _number(_value(r, "days_until_expiry"))) is not None and days < 0
        for r in unique_keys.values()
    )
    _section(sheet, 11, "Issue Summary", RED)
    _block(sheet, 12,
           f"{aged} credentials are reported active, linked to highly privileged identities, "
           "and last rotated more than 365 days ago.", fill="FFE8E8")
    if expiring:
        earliest = min(expiring, key=lambda r: int(r["days_until_expiry"]))
        expiry_message = (
            f"{len(expiring)} credentials have reported expiry dates within 60 days; "
            f"the earliest is {_value(earliest, 'valid_before')}. "
        )
    else:
        expiry_message = "No supplied expiry date falls within the next 60 days. "
    expiry_message += (
        f"{expired} have reported expiry dates in the past."
    )
    _block(sheet, 13, expiry_message, fill="FFE8E8")
    sheet["A13"].comment = _comment("Expiry dates are not remediation deadlines.")
    _block(sheet, 14,
           "These credentials authenticate AWS or Azure identities and allow access to "
           "resources within those identities' effective permissions.", fill="FFE8E8")
    sheet["A14"].comment = _comment(
        "Missing information is marked as unknown. An inactivity flag alone is not a reason "
        "to remove a credential."
    )
    _section(sheet, 16, "Risk Description", ORANGE)
    _block(sheet, 17,
           "A compromised credential could enable unauthorized changes within its effective "
           "permissions. Broad IAM or directory actions do not prove unrestricted cloud access. "
           "This finding is not evidence of a compromise or data loss.", fill="FFF4E6")
    _section(sheet, 19, "Security Recommendation and Requested Response", GREEN)
    _block(sheet, 20,
           "- We recommend staged replacement of required credentials and controlled "
           "retirement of credentials confirmed as obsolete.\n"
           "- Changes are subject to dependency testing, a recovery plan and change approval.",
           fill="E8F5E8")
    _block(sheet, 21,
           "- Please confirm the preferred approach and any business constraints.\n"
           "- Please confirm ownership, required access and application dependencies.",
           fill="E8F5E8", size=10)
    severities = _counts([_severity(_value(r, "issue_severity")) for r in unique_issues.values()])
    statuses = _counts([_value(r, "issue_status") for r in unique_issues.values()])
    if unique_issues or issues is not None:
        controls = ", ".join(sorted({
            _value(r, "issue_control_id") for r in unique_issues.values()
        })) or "Unknown"
        metadata = (
            f"Wiz issues: {len(unique_issues)} | "
            f"Wiz severity: {severities or 'Unknown'} | "
            f"Wiz status: {statuses or 'Unknown'} | Control ID: {controls}."
        )
        context = "Source issue metadata is derived from matched issue CSV records.\n"
    else:
        metadata = (
            f"Provided context (not verified): severity {_severity(severity)} | "
            f"Control ID: {rule_id or 'Unknown'}. Source issue severity: Unknown."
        )
        context = "Provided context is not verified source issue metadata.\n"
    _metric(sheet, 23, metadata + "\nRemediation approval: Pending.")
    context += (
        f"Distinct source issue lifecycle statuses: {statuses or 'Unknown'}\n"
        "Source CSV files remain the original evidence; this report is not a raw archive.\n"
        "Issue updated-at timestamps describe issue lifecycle, not fresh credential observations.\n"
        "No approval, business risk acceptance, execution or verified closure is asserted."
    )
    sheet["A23"].comment = _comment(context)
    if workflow_text or workflow_context:
        sheet["A21"].comment = _comment(
            "\n".join(part for part in (workflow_text, workflow_context) if part)
        )
    sheet.print_area = "A1:A23"
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.tabColor = NAVY


def _display(record: Mapping[str, str], *, enriched: bool) -> dict[str, str]:
    account = _value(record, "scope_name", _value(record, "scope_external_id"))
    native_id = _value(record, "scope_external_id", "")
    if native_id and native_id != account:
        account += "\n" + native_id
    if "tenant" in _value(record, "scope_type", "").casefold():
        account = "Entra tenant\n" + account
    return {
        "account": account,
        "principal": _value(record, "principal", _value(record, "issue_resource_name")),
        "credential": _value(record, "credential_name", _value(record, "credential_id")),
        "rotated": _value(record, "rotated_date", _value(record, "rotated_at")),
        "age": _value(record, "age_days"),
        "expiry": _value(record, "expiry_date", _value(record, "valid_before", "Not supplied")),
        "permissions": _value(record, "permission_names"),
        "inactivity": {
            "true": "No activity in 90 days (Wiz)",
            "false": "Not flagged inactive (Wiz)",
        }.get(_value(record, "inactive").casefold(), "Unknown"),
        "status": (
            f"{_value(record, 'issue_status')}\nApproval not recorded" if enriched
            else "Pending review\nApproval not recorded"
        ),
    }


def _issue_url(record: Mapping[str, str]) -> str | None:
    url = _value(record, "issue_url", "")
    issue_id = _value(record, "issue_id", "")
    if not url:
        return None
    if not issue_id:
        raise ValueError("A nonblank Wiz URL requires a source issue ID.")
    return validate_wiz_url(url, issue_id, location="report row") or None


def resource_table(
    workbook: Workbook, name: str, records: Sequence[Mapping[str, str]],
) -> None:
    if len(records) > MAX_ROWS - 1:
        raise ValueError("Report exceeds the Excel 1048576-row worksheet limit.")
    sheet = workbook.create_sheet(name)
    _print_setup(sheet, landscape=True)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.selection = [
        Selection(pane="bottomLeft", activeCell="A2", sqref="A2"),
    ]
    sheet.sheet_properties.tabColor = BLUE if name == "AWS Data" else "2E75B6"
    sheet.print_title_rows = "1:1"
    enriched = any(any(key.startswith("issue_") for key in r) for r in records)
    for column, (_, label, width) in enumerate(TABLE_FIELDS, 1):
        if column == 9 and enriched:
            label = "Wiz Status / Approval"
        _literal(sheet, 1, column, label)
        cell = sheet.cell(1, column)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 42
    sorted_records = sorted(
        records,
        key=lambda r: tuple(_value(r, key).casefold() for key in (
            "cloud", "scope_name", "principal", "credential_name", "credential_id",
        )),
    )
    for row, record in enumerate(sorted_records, 2):
        _identity(record)
        _credential(record)
        values = _display(record, enriched=enriched)
        height = max(_height(values[key], width, minimum=40) for key, _, width in TABLE_FIELDS)
        for column, (key, _, _) in enumerate(TABLE_FIELDS, 1):
            _literal(sheet, row, column, values[key])
            cell = sheet.cell(row, column)
            cell.font = Font(name="Calibri", size=10, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=STRIPE if row % 2 == 0 else "FFFFFF")
            cell.border = THIN_BORDER
        context = "\n".join(
            f"{label}: {_value(record, key)}" for key, label in COMMENT_FIELDS
            if key in record
        )
        context += (
            "\nSource CSV files remain the original evidence."
            "\nWiz lifecycle status is not approval, risk acceptance or verified closure."
            "\nThis workbook contains read-only proposals; approval is not recorded."
        )
        sheet.cell(row, 2).comment = _comment(context)
        if url := _issue_url(record):
            sheet.cell(row, 2).hyperlink = url
        sheet.row_dimensions[row].height = height
    last = max(1, len(records) + 1)
    sheet.auto_filter.ref = f"A1:I{last}"
    sheet.print_area = f"A1:I{last}"


def remediation_page(workbook: Workbook, *, notes: str = "") -> None:
    _check_text(notes, notes=True)
    sheet = _page(workbook, "Remediation")
    _title(sheet, "Remediation Approach", "Proposed Credential Remediation")
    if notes:
        label = "Provided review notes - not validated and not approval:\n"
        _check_text(label + notes, notes=True)
        sheet["A1"].comment = _comment(label + notes)
    _section(sheet, 5, "Approval Required Before Changes", ORANGE)
    _block(sheet, 6,
           "Agree the affected resources, change window, testing and recovery plan with the "
           "owner. Obtain approval before making changes. A Wiz issue status does not replace "
           "change approval, risk acceptance or confirmation that the work is complete.",
           fill="FFF4E6")
    steps = (
        ("Step 1: Confirm Ownership and Usage",
         "Identify the application, identity and credential owners. Verify native credential "
         "IDs, type, effective permissions, last use and all consumers. A Wiz inactivity flag "
         "alone does not justify removal. Check approved lifecycle standards and exceptions."),
        ("Step 2: If No Longer Needed",
         "Request controlled retirement of the confirmed-obsolete credential. Agree on impact "
         "checks, observation and recovery first. Do not delete an identity or credential "
         "solely because it looks inactive; infrequent jobs may still depend on it."),
        ("Step 3: If Still Required - AWS",
         "Plan an approved overlapping replacement. Check the two-access-key limit and both "
         "keys' consumers first. Securely update applications, validate business cycles, then "
         "deactivate the old key before agreed deletion. Never disclose secret values in tickets."),
        ("Step 4: If Still Required - Azure / Entra",
         "Confirm secret versus certificate and the application/service principal owning it. "
         "Use its supported rollover process, update every consumer and test authentication "
         "before retiring the old credential. Removed credentials may not be recoverable."),
        ("Step 5: Reduce Unnecessary Access",
         "Validate broad permissions against actual workload needs and guardrails. Pilot "
         "least-privilege changes separately where safer. Consider roles, managed identities "
         "or federation only where supported by the existing architecture and vendor."),
        ("Step 6: Validate, Recover and Close",
         "Record non-secret before/after evidence and application-owner checks. Stop on "
         "unexpected impact. Do not reactivate suspected-compromised credentials as rollback. "
         "Obtain a fresh Wiz assessment; suppression or risk acceptance is not remediation."),
        ("Expected Outcome",
         "Required credentials are governed by an approved lifecycle and appropriate access; "
         "obsolete credentials are retired safely. Ownership, approvals and validation are "
         "documented. No technical change or acceptance is asserted by this workbook."),
    )
    row = 8
    for index, (heading, detail) in enumerate(steps):
        if index == 4:
            sheet.row_breaks.append(Break(id=row - 1))
        row = _section(sheet, row, heading)
        row = _block(sheet, row, detail, fill="E8F5E8")
        sheet.row_dimensions[row].height = 10
        row += 1
    sheet.print_area = f"A1:A{row - 1}"
    sheet.sheet_properties.tabColor = GREEN
