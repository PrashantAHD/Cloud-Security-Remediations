"""Synthetic progress tracking and report snapshot tests."""

import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_csv_report import export, record, workspace  # noqa: F401

from remediation import csv_report, workflow
from remediation.__main__ import main
from remediation.workflow import GATES, create_tracker, load_tracker, record_progress


def create(path):
    return create_tracker(
        path, finding_id="synthetic-case", control_id="synthetic-control",
        scope="Synthetic project; open findings; snapshot date supplied separately",
        summary="Sample page collected; graph collection remains outstanding.",
        next_action="Collect and reconcile both sources.",
    )


def progress(path, stage=1, status="in_progress", **kwargs):
    return record_progress(
        path, stage=stage, status=status, summary="Reviewed synthetic evidence.",
        next_action="Perform the next reviewed action.", **kwargs,
    )


def test_init_persist_show_and_no_overwrite(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    tracker = create(path)
    assert load_tracker(path) == tracker
    assert tracker.current_stage == 1
    assert tracker.status == "in_progress"
    assert tracker.stage_status(2) == "not_started"
    assert "graph collection remains outstanding" in tracker.display()
    assert "Draft - evidence review incomplete" in tracker.cover_text()
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        create(path)
    assert path.read_bytes() == before
    assert not path.with_suffix(".json.lock").exists()
    assert main(["workflow", "show", "--tracker", str(path)]) == 0


def test_full_sequential_workflow_keeps_history(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    for stage in range(1, 9):
        tracker = progress(
            path, stage, "complete",
            evidence=dict.fromkeys(GATES[stage - 1], "Reviewed private evidence reference"),
        )
        assert tracker.current_stage == min(8, stage + 1)
        assert tracker.status == ("complete" if stage == 8 else "in_progress")
        assert len(tracker.events) == stage + 1
    assert all(tracker.stage_status(n) == "complete" for n in range(1, 9))
    assert load_tracker(path) == tracker
    with pytest.raises(ValueError, match="cannot be skipped or reopened"):
        progress(path, 8)


@pytest.mark.parametrize("kwargs", [
    {"stage": 2},
    {"stage": True},
    {"status": "complete"},
    {"status": "complete", "evidence": {"issues": "Only a sample"}},
    {"status": "complete", "blocker": "Missing graph"},
    {"status": "blocked"},
    {"status": "in_progress", "blocker": "Missing graph"},
    {"evidence": {"unknown": "Invalid gate"}},
    {"evidence": {"issues": ""}},
    {"status": "invalid"},
])
def test_invalid_changes_leave_bytes_unchanged(workspace, kwargs):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    before = path.read_bytes()
    with pytest.raises(ValueError):
        progress(path, **kwargs)
    assert path.read_bytes() == before


def test_blocked_and_resumed_with_explicit_update(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    blocked = progress(path, status="blocked", blocker="Source unavailable")
    assert blocked.status == "blocked"
    assert "Source unavailable" in blocked.display()
    resumed = progress(path)
    assert resumed.status == "in_progress"
    assert resumed.events[1].blocker == "Source unavailable"
    assert resumed.events[-1].blocker == ""


@pytest.mark.parametrize("mutate", [
    lambda d: d.update(version=2),
    lambda d: d.update(version=True),
    lambda d: d.update(events=[]),
    lambda d: d.update(events="invalid"),
    lambda d: d.update(scope=""),
    lambda d: d.update(unexpected=True),
    lambda d: d["events"][0].update(stage=4),
    lambda d: d["events"][0].update(status="complete"),
    lambda d: d["events"][0].update(at="not a date"),
    lambda d: d["events"][0].update(at="2026-01-01T00:00:00"),
    lambda d: d["events"][0].update(evidence=[]),
    lambda d: d["events"][0].update(summary="unsafe\ntext"),
])
def test_corrupt_tracker_rejected(workspace, mutate):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    data = json.loads(path.read_text())
    mutate(data)
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError):
        load_tracker(path)


def test_duplicate_keys_and_size_limit():
    with pytest.raises(ValueError, match="duplicate"):
        workflow._decode(b'{"version": 1, "version": 1}')
    with pytest.raises(ValueError, match="size limit"):
        workflow._decode(b" " * (workflow.MAX_BYTES + 1))


def test_cli_completion_and_duplicate_reference_errors(workspace, capsys):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    args = [
        "workflow", "record", "--tracker", str(path), "--stage", "1",
        "--status", "complete", "--summary", "Reviewed full synthetic collection.",
        "--next-action", "Analyze validated findings.",
    ]
    for gate in GATES[0]:
        args.extend(["--evidence", f"{gate}=Reviewed private reference"])
    before = path.read_bytes()
    assert main([*args, "--evidence", "issues=Duplicate"]) == 1
    assert "unique GATE=REFERENCE" in capsys.readouterr().err
    assert path.read_bytes() == before
    assert main(args) == 0
    assert load_tracker(path).current_stage == 2
    assert "Stage 2/8" in capsys.readouterr().out


def test_cli_missing_tracker_fails_without_creation(workspace, capsys):  # noqa: F811
    path = workspace / "missing.json"
    assert main(["workflow", "show", "--tracker", str(path)]) == 1
    assert "cannot read or publish" in capsys.readouterr().err
    assert not path.exists()


def test_symlink_tracker_is_rejected(workspace, monkeypatch):  # noqa: F811
    path = workspace / "tracker.json"
    monkeypatch.setattr(Path, "is_symlink", lambda self: True)
    with pytest.raises(ValueError, match="non-symlink"):
        create(path)


@pytest.mark.parametrize("path", [Path("relative.json"), workflow.ROOT / "private.json"])
def test_private_location_required(path):
    with pytest.raises(ValueError):
        create(path)


def test_lock_and_concurrent_change_preserve_tracker(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    tracker = create(path)
    before = path.read_bytes()
    lock = path.with_suffix(".json.lock")
    lock.write_text("existing writer", encoding="utf-8")
    with pytest.raises(FileExistsError):
        progress(path)
    assert lock.read_text() == "existing writer"
    lock.unlink()
    with pytest.raises(ValueError, match="changed"):
        workflow._save(path, tracker, b"stale digest")
    assert path.read_bytes() == before
    assert not lock.exists()


def test_failed_publish_cleans_own_lock(workspace, monkeypatch):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    before = path.read_bytes()

    def failed(*args, **kwargs):
        raise PermissionError("Synthetic locked destination")

    monkeypatch.setattr(csv_report, "_publish", failed)
    with pytest.raises(PermissionError):
        progress(path)
    assert path.read_bytes() == before
    assert not path.with_suffix(".json.lock").exists()


def test_timestamp_order_is_validated(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    tracker = create(path)
    old = replace(tracker.events[0], at="2000-01-01T00:00:00+00:00")
    with pytest.raises(ValueError, match="out of order"):
        workflow._validate(replace(tracker, events=(*tracker.events, old)))


def test_cli_records_errors_and_literal_cover_snapshot(workspace, capsys):  # noqa: F811
    path = workspace / "tracker.json"
    assert main([
        "workflow", "init", "--tracker", str(path),
        "--finding-id", "=synthetic-case", "--control-id", "synthetic-control",
        "--scope", "Synthetic scope", "--summary", "Sample only",
        "--next-action", "Fetch graph",
    ]) == 0
    assert main([
        "workflow", "record", "--tracker", str(path), "--stage", "1",
        "--status", "complete", "--summary", "Not sufficient",
        "--next-action", "Analyze",
    ]) == 1
    assert "Completion requires" in capsys.readouterr().err
    source = export(workspace, [record()])
    report = workspace / "report.xlsx"
    before = path.read_bytes()
    for entrypoint in (main, csv_report.main):
        args = [str(source), "--output", str(report), "--workflow", str(path),
                "--rule-id", "synthetic-control"]
        if report.exists():
            args.append("--update")
        assert entrypoint(args) == 0
    book = load_workbook(report)
    try:
        assert book.sheetnames == csv_report.SHEETS
        assert "Please confirm the preferred approach" in book["Cover"]["A21"].value
        assert load_tracker(path).cover_text() in book["Cover"]["A21"].comment.text
        visible = "\n".join(str(c.value) for row in book["Cover"] for c in row if c.value)
        for internal_text in ("Stage 1/8", "Fetch graph", "=synthetic-case", "Synthetic scope"):
            assert internal_text not in visible
        assert book["Cover"]["A21"].data_type == "s"
        assert "Finding: =synthetic-case" in book["Cover"]["A21"].comment.text
        assert "Scope: Synthetic scope" in book["Cover"]["A21"].comment.text
        assert book["Cover"].page_setup.fitToHeight == 1
        assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    finally:
        book.close()
    assert path.read_bytes() == before
    with pytest.raises(ValueError, match="control ID must match"):
        csv_report.generate_report(
            [source], output=workspace / "wrong.xlsx", workflow=path, rule_id="different",
        )
    assert not (workspace / "wrong.xlsx").exists()


def test_report_update_refreshes_snapshot_without_advancing_tracker(workspace):  # noqa: F811
    path = workspace / "tracker.json"
    create(path)
    source = export(workspace, [record()])
    report = workspace / "report.xlsx"
    options = {"output": report, "workflow": path, "rule_id": "synthetic-control"}
    csv_report.generate_report([source], **options)
    progress(path, status="blocked", blocker="Missing graph")
    csv_report.generate_report([source], update=True, **options)
    book = load_workbook(report)
    try:
        assert "Status: blocked" in book["Cover"]["A21"].comment.text
        assert "Missing graph" in book["Cover"]["A21"].comment.text
        visible = "\n".join(str(c.value) for row in book["Cover"] for c in row if c.value)
        assert "Status: blocked" not in visible and "Missing graph" not in visible
    finally:
        book.close()
    assert load_tracker(path).current_stage == 1
    del options["workflow"]
    csv_report.generate_report([source], update=True, **options)
    book = load_workbook(report)
    try:
        assert "Please confirm the preferred approach" in book["Cover"]["A21"].value
        assert book["Cover"]["A21"].comment is None
    finally:
        book.close()
