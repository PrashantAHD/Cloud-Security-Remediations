"""Synthetic graph CSV tests; fixtures never contain real cloud evidence."""

import csv
import os
import shutil
import subprocess
import sys
import time
from copy import copy
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from remediation import csv_report
from remediation.csv_report import (
    KEY,
    PERMISSION,
    PRINCIPAL,
    REPOSITORY_ROOT,
    SCOPE,
    generate_report,
    main,
)


def retry_locked_cleanup(function, path, error):
    """Retry only transient Windows sharing/lock violations during test cleanup."""
    for delay in (0.1, 0.2, 0.4, 0.8, 1.6):
        if not isinstance(error, PermissionError) or getattr(error, "winerror", None) not in (
            32, 33,
        ):
            raise error
        time.sleep(delay)
        try:
            function(path)
            return
        except OSError as retry_error:
            error = retry_error
    raise error


@pytest.fixture
def workspace():
    # Output must be outside the repository; use its workspace parent, not system temp.
    root = REPOSITORY_ROOT.parent / f".csv-report-test-{uuid4().hex}"
    root.mkdir()
    try:
        yield root
    finally:
        shutil.rmtree(root, onexc=retry_locked_cleanup)


def test_cleanup_retries_transient_windows_lock(monkeypatch, tmp_path):
    path = tmp_path / "synthetic.txt"
    path.write_text("synthetic", encoding="utf-8")
    delays = []
    calls = []
    monkeypatch.setattr(time, "sleep", delays.append)
    error = PermissionError("synthetic sharing violation")
    error.winerror = 32

    def unlink(value):
        calls.append(value)
        if len(calls) == 1:
            raise error
        value.unlink()

    retry_locked_cleanup(unlink, path, error)
    assert calls == [path, path]
    assert delays == [0.1, 0.2]
    assert not path.exists()


@pytest.mark.parametrize("winerror", [32, 33])
def test_cleanup_persistent_lock_still_fails(monkeypatch, winerror):
    error = PermissionError("synthetic persistent lock")
    error.winerror = winerror
    delays = []
    monkeypatch.setattr(time, "sleep", delays.append)

    def locked(_):
        raise error

    with pytest.raises(PermissionError) as caught:
        retry_locked_cleanup(locked, "synthetic-path", error)
    assert caught.value is error
    assert delays == [0.1, 0.2, 0.4, 0.8, 1.6]


@pytest.mark.parametrize("error", [PermissionError("denied"), OSError("unrelated failure")])
def test_cleanup_does_not_retry_unrelated_errors(monkeypatch, error):
    def unexpected(_):
        pytest.fail("Unrelated cleanup errors must not be retried.")

    monkeypatch.setattr(time, "sleep", unexpected)
    with pytest.raises(type(error)) as caught:
        retry_locked_cleanup(unexpected, "synthetic-path", error)
    assert caught.value is error


def test_cleanup_stops_if_retry_error_changes(monkeypatch):
    error = PermissionError("synthetic sharing violation")
    error.winerror = 32
    failure = OSError("synthetic different failure")
    delays = []
    monkeypatch.setattr(time, "sleep", delays.append)

    def changed(_):
        raise failure

    with pytest.raises(OSError) as caught:
        retry_locked_cleanup(changed, "synthetic-path", error)
    assert caught.value is failure
    assert delays == [0.1]


def record(**overrides):
    row = {
        PRINCIPAL + "ID": "principal-one",
        PRINCIPAL + "Name": "Synthetic principal",
        PRINCIPAL + "RowObjectType": "SERVICE_ACCOUNT",
        PRINCIPAL + "displayName": "Synthetic display name",
        PRINCIPAL + "hasHighPrivileges": "true",
        PRINCIPAL + "hasMfa": "",
        PRINCIPAL + "inactiveInLast90Days": "false",
        PRINCIPAL + "providerUniqueId": "native-principal",
        PRINCIPAL + "status": "source-state",
        KEY + "ID": "graph-key-one",
        KEY + "Name": "Synthetic credential display",
        KEY + "active": "true",
        KEY + "rotatedAt": "2026-01-01T00:00:00Z",
        KEY + "validAfter": "",
        KEY + "validBefore": "2026-04-01T00:00:00Z",
        PERMISSION + "ID": "permission-one",
        PERMISSION + "Name": "Permission one",
        PERMISSION + "accessTypes": "READ",
        PERMISSION + "cloudPlatform": "AWS",
        PERMISSION + "status": "source-state",
        SCOPE + "ID": "scope-one",
        SCOPE + "Name": "Synthetic scope",
        SCOPE + "RowObjectType": "SUBSCRIPTION",
        SCOPE + "cloudPlatform": "AWS",
        SCOPE + "encryptsSecrets": "",
        SCOPE + "externalId": "0000000042",
        SCOPE + "isManaged": "",
        SCOPE + "isPaaS": "",
        SCOPE + "nativeType": "SyntheticScope",
        SCOPE + "status": "source-state",
        SCOPE + "subscriptionId": "0000000042",
    }
    row.update(overrides)
    return row


def export(workspace, rows, name="input.csv", headers=None):
    path = workspace / name
    fields = list(rows[0]) if headers is None else headers
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(fields)
        writer.writerows([row.get(field, "") for field in fields] for row in rows)
    return path


def report(workspace, rows, **kwargs):
    source = export(workspace, rows)
    result = generate_report(
        [source], output=workspace / "report.xlsx", as_of="2026-04-01", **kwargs
    )
    return result, load_workbook(result.path)


def sheet_rows(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def test_group_four_permissions_two_keys_multiple_exports(workspace):
    rows = [
        record(**{
            KEY + "ID": key,
            PERMISSION + "ID": f"permission-{number}",
            PERMISSION + "Name": f"Permission {number}",
            PERMISSION + "accessTypes": "READ" if number % 2 else "WRITE",
        })
        for key in ("graph-key-one", "graph-key-two")
        for number in range(4)
    ]
    first = export(workspace, rows)
    second_folder = workspace / "second"
    second_folder.mkdir()
    second = export(second_folder, rows)
    result = generate_report(
        [first, second], output=workspace / "report.xlsx", as_of="2026-04-01"
    )
    assert (result.source_rows, result.unique_principals) == (16, 1)
    assert (result.unique_credentials, result.grouped_scopes) == (2, 2)
    assert (result.aws_principals, result.aws_credentials) == (1, 2)
    assert (result.azure_principals, result.azure_credentials) == (0, 0)
    workbook = load_workbook(result.path)
    assert workbook.sheetnames == csv_report.SHEETS
    evidence = csv_report._read_exports([first, second], date(2026, 4, 1))
    assert {row["source_file"] for row in evidence.source} == {"input.csv"}
    assert {row["source_input"] for row in evidence.source} == {"1", "2"}
    grouped = evidence.records
    assert len(grouped) == 2
    for row in grouped:
        assert row["permission_names"].splitlines() == [
            f"Permission {number}" for number in range(4)
        ]
        assert set(row["permission_access_types"].splitlines()) == {"READ", "WRITE"}
        assert len(row["source_rows"].splitlines()) == 8
        assert row["status"] == "Draft"
    assert workbook["AWS Data"].max_row == 3
    assert workbook["Azure Data"].max_row == 1
    assert result.unique_issues is None
    assert workbook["AWS Data"]["G2"].value == grouped[0]["permission_names"]


def test_cloud_and_scope_collisions_do_not_merge(workspace):
    rows = [
        record(),
        record(**{SCOPE + "ID": "scope-two"}),
        record(**{SCOPE + "cloudPlatform": "Azure", PERMISSION + "cloudPlatform": "Azure"}),
    ]
    result, workbook = report(workspace, rows)
    assert (result.unique_principals, result.unique_credentials, result.grouped_scopes) == (2, 2, 3)
    assert (result.aws_principals, result.azure_principals) == (1, 1)
    assert workbook["AWS Data"].max_row == 3
    assert workbook["Azure Data"].max_row == 2


def test_permission_access_type_arrays_merge_unique_members(workspace):
    _, workbook = report(workspace, [
        record(**{PERMISSION + "accessTypes": '["READ", "WRITE", "READ"]'}),
        record(**{PERMISSION + "accessTypes": '["WRITE", "ADMIN"]'}),
    ])
    evidence = csv_report._read_exports([workspace / "input.csv"], date(2026, 4, 1))
    assert evidence.records[0]["permission_access_types"].splitlines() == ["READ", "WRITE", "ADMIN"]
    assert "ADMIN" in workbook["AWS Data"]["B2"].comment.text


def test_azure_scope_type_preserves_tenant_evidence_and_unknown_key_type(workspace):
    _, workbook = report(workspace, [record(**{
        SCOPE + "cloudPlatform": "Azure",
        PERMISSION + "cloudPlatform": "",
        SCOPE + "RowObjectType": "SUBSCRIPTION",
        SCOPE + "nativeType": "tenant",
        SCOPE + "subscriptionId": "",
    })])
    evidence = csv_report._read_exports([workspace / "input.csv"], date(2026, 4, 1))
    row = evidence.records[0]
    assert row["scope_type"] == (
        "RowObjectType: SUBSCRIPTION; Native type: tenant; Subscription ID: Unknown"
    )
    assert row["scope_type"] in workbook["Azure Data"]["B2"].comment.text
    assert "credential_type" not in row
    assert "mfa" not in row


@pytest.mark.parametrize(
    ("rotated", "expiry", "age", "remaining"),
    [
        ("2026-04-01", "2026-04-01", "0", "0"),
        ("2026-03-31T23:30:00-02:00", "2026-04-02T00:30:00+02:00", "0", "0"),
        ("2026-04-01T00:30:00+02:00", "2026-03-31", "1", "-1"),
        ("2026-01-01", "2026-04-02", "90", "1"),
        ("2026-03-31", "", "1", "Not supplied"),
    ],
)
def test_utc_age_and_expiry_boundaries(workspace, rotated, expiry, age, remaining):
    _, workbook = report(workspace, [record(**{
        KEY + "rotatedAt": rotated, KEY + "validBefore": expiry,
    })])
    row = csv_report._read_exports([workspace / "input.csv"], date(2026, 4, 1)).records[0]
    assert row["age_days"] == age
    assert row["days_until_expiry"] == remaining
    assert str(workbook["AWS Data"]["E2"].value) == age
    expected_rotated = csv_report._utc_date(rotated, "rotatedAt").isoformat()
    expected_expiry = csv_report._utc_date(expiry, "validBefore")
    assert workbook["AWS Data"]["D2"].value == expected_rotated
    assert workbook["AWS Data"]["F2"].value == (
        expected_expiry.isoformat() if expected_expiry else "Not supplied"
    )
    assert rotated in workbook["AWS Data"]["B2"].comment.text


def test_missing_dates_and_flags_are_not_inferred(workspace):
    _, workbook = report(workspace, [record(**{
        KEY + "rotatedAt": " ",
        KEY + "validBefore": "",
        KEY + "active": "",
        PRINCIPAL + "hasHighPrivileges": " ",
        PRINCIPAL + "inactiveInLast90Days": "",
        PRINCIPAL + "RowObjectType": "",
    })])
    row = csv_report._read_exports([workspace / "input.csv"], date(2026, 4, 1)).records[0]
    for field in (
        "rotated_at", "age_days", "key_active", "high_privileges",
        "inactive", "principal_type",
    ):
        assert row[field] == "Unknown"
    assert row["valid_before"] == "Not supplied"
    assert row["days_until_expiry"] == "Not supplied"
    assert workbook["AWS Data"]["D2"].value == "Unknown"


def test_literal_values_and_original_sources_stay_unchanged(workspace):
    long_value = "=1+1\n" + "Synthetic evidence " * 4000
    payload = '=HYPERLINK("https://example.invalid/","Synthetic")'
    first = record(**{
        PRINCIPAL + "Name": payload,
        KEY + "Name": "+1+1",
        SCOPE + "Name": "https://example.invalid/",
        "=EXTRA_HEADER()": long_value,
        "Other field": "@SUM(1,2)",
    })
    source = export(workspace, [first, first])
    other = record(**{
        KEY + "ID": "key-other", "New field": "-1+2",
        PRINCIPAL + "ID": "principal-other", SCOPE + "ID": "scope-other",
    })
    second = export(workspace, [other], "second.csv")
    notes = workspace / "notes.txt"
    note_text = "=Synthetic unverified note\n+Do not interpret as approval"
    notes.write_text(note_text, encoding="utf-8")
    original = source.read_bytes()
    result = generate_report(
        [source, second], output=workspace / "report.xlsx", as_of="2026-04-01",
        title=payload, rule_id="=1+1", severity="High", notes=notes,
    )
    workbook = load_workbook(result.path)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                assert cell.hyperlink is None
                if cell.value is not None:
                    assert cell.data_type == "s"
    data = sheet_rows(workbook["AWS Data"])[0]
    assert "0000000042" in data["Account / Scope"]
    assert data["Resource Name"] == payload
    assert "+1+1" in data["Credential"]
    evidence = csv_report._read_exports([source, second], date(2026, 4, 1))
    raw = evidence.source
    assert len(raw) == 3
    source_fields = {label: key for key, label in evidence.source_fields}
    for field, value in first.items():
        assert raw[0][source_fields[field]] == value
    assert source.read_bytes() == original
    assert raw[0]["source_file"] == "input.csv"
    assert raw[0]["source_input"] == "1"
    assert raw[0]["source_row"] == "2"
    assert int(raw[1]["source_row"]) > 2
    assert raw[2][source_fields["New field"]] == "-1+2"
    assert raw[2]["source_input"] == "2"
    assert note_text in workbook["Remediation"]["A1"].comment.text


def test_notes_are_retained_without_extra_tabs(workspace):
    notes = (
        "=Synthetic heading\n\n"
        + "A synthetic assessment paragraph with readable words. " * 230
        + "\n\n+Literal finding\n@Literal follow-up\n\tIndented observation\n"
        + "x" * 1100
    )
    note_path = workspace / "notes.txt"
    note_path.write_text(notes, encoding="utf-8")
    _, workbook = report(workspace, [record()], notes=note_path)
    assert workbook.sheetnames == csv_report.SHEETS
    assert notes in workbook["Remediation"]["A1"].comment.text


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        (PRINCIPAL + "ID", "", "grouping IDs"),
        (KEY + "ID", "  ", "grouping IDs"),
        (SCOPE + "ID", "", "grouping IDs"),
        (SCOPE + "cloudPlatform", "OtherCloud", "Unsupported cloud"),
        (PERMISSION + "cloudPlatform", "Azure", "conflict"),
        (KEY + "rotatedAt", "yesterday", "rotatedAt"),
        (KEY + "rotatedAt", "2026-02-30", "rotatedAt"),
        (KEY + "rotatedAt", "2026-04-01T00:00:00", "rotatedAt"),
        (KEY + "rotatedAt", "2026-04-02", "after"),
        (KEY + "validBefore", "invalid-expiry", "validBefore"),
        (KEY + "validBefore", "2026-13-01", "validBefore"),
        (PERMISSION + "accessTypes", '["READ", 3]', "text values"),
        (PERMISSION + "accessTypes", "[unclosed", "invalid JSON array"),
        (PRINCIPAL + "Name", "synthetic\x01text", "characters"),
    ],
)
def test_invalid_evidence_is_rejected_without_partial_output(workspace, field, value, message):
    source = export(workspace, [record(**{field: value})])
    with pytest.raises(ValueError, match=message):
        generate_report([source], output=workspace / "report.xlsx", as_of="2026-04-01")
    assert not (workspace / "report.xlsx").exists()
    assert not list(workspace.glob("*.partial"))


def test_blank_cloud_rejected_and_permission_cloud_fallback_accepted(workspace):
    source = export(workspace, [record(**{
        SCOPE + "cloudPlatform": "", PERMISSION + "cloudPlatform": "",
    })])
    with pytest.raises(ValueError, match="Cloud must be supplied"):
        generate_report([source], output=workspace / "report.xlsx")
    result, _ = report(workspace, [record(**{SCOPE + "cloudPlatform": ""})])
    assert result.aws_credentials == 1


def test_drift_is_explicit_not_arbitrarily_selected(workspace):
    source = export(workspace, [record(), record(**{KEY + "active": "false"})])
    with pytest.raises(ValueError, match="input 1, row 2 and input 1, row 3;.*snapshot drift"):
        generate_report([source], output=workspace / "report.xlsx", as_of="2026-04-01")


@pytest.mark.parametrize(
    ("field", "value"),
    [
        (PRINCIPAL + "displayName", "Different synthetic display name"),
        (PRINCIPAL + "hasMfa", "true"),
        (SCOPE + "encryptsSecrets", "true"),
        (KEY + "validAfter", "2026-02-01"),
        (KEY + "rotatedAt", "2026-02-01"),
        (KEY + "active", "false"),
    ],
)
def test_cross_export_drift_including_non_displayed_facts_has_safe_row_references(
    workspace, field, value
):
    first_dir = workspace / "first"
    second_dir = workspace / "second"
    first_dir.mkdir()
    second_dir.mkdir()
    first = export(first_dir, [record()])
    second = export(second_dir, [record(**{field: value})])
    with pytest.raises(
        ValueError, match="input 1, row 2 and input 2, row 2;.*snapshot drift"
    ) as error:
        generate_report([first, second], output=workspace / "report.xlsx", as_of="2026-04-01")
    assert "Synthetic" not in str(error.value)
    assert not (workspace / "report.xlsx").exists()


@pytest.mark.parametrize(
    "case", ["empty", "headers", "duplicate", "blank", "missing", "cloud", "ragged"]
)
def test_invalid_csv_shapes(workspace, case):
    row = record()
    fields = list(row)
    source = workspace / "bad.csv"
    if case == "empty":
        source.write_text("", encoding="utf-8")
    elif case == "headers":
        source = export(workspace, [], headers=fields)
    elif case == "duplicate":
        source = export(workspace, [row], headers=[*fields, fields[0]])
    elif case == "blank":
        source = export(workspace, [row], headers=[*fields, " "])
    elif case == "missing":
        source = export(workspace, [row], headers=fields[1:])
    elif case == "cloud":
        source = export(workspace, [row], headers=[
            field for field in fields if not field.endswith(".cloudPlatform")
        ])
    else:
        source = export(workspace, [row])
        with source.open("a", encoding="utf-8") as handle:
            handle.write("ragged,row\n")
    with pytest.raises(ValueError):
        generate_report([source], output=workspace / "report.xlsx", as_of="2026-04-01")


@pytest.mark.parametrize("content", [b"\xff", b'"unterminated'])
def test_invalid_utf8_and_csv_quoting_cli_errors_are_sanitized(workspace, content, capsys):
    source = workspace / "sensitive-name.csv"
    source.write_bytes(content)
    assert main([str(source), "--output", str(workspace / "report.xlsx")]) == 1
    captured = capsys.readouterr()
    assert not captured.out
    assert "sensitive-name" not in captured.err


def test_output_protection_and_no_overwrite(workspace):
    source = export(workspace, [record()])
    destination = workspace / "report.xlsx"
    destination.write_bytes(b"existing bytes")
    for output, message in (
        (destination, "overwrite"),
        (REPOSITORY_ROOT / "must-not-write.xlsx", "outside"),
        (Path("relative.xlsx"), "absolute"),
        (workspace / "wrong.txt", ".xlsx"),
    ):
        with pytest.raises(ValueError, match=message):
            generate_report([source], output=output)
    assert destination.read_bytes() == b"existing bytes"


def test_atomic_publication_refuses_racing_destination(workspace, monkeypatch):
    output = workspace / "report.xlsx"
    link = os.link

    def racing_link(source, destination):
        Path(destination).write_bytes(b"concurrent writer")
        link(source, destination)

    monkeypatch.setattr(csv_report.os, "link", racing_link)
    with pytest.raises(FileExistsError):
        csv_report._publish(BytesIO(b"complete report"), output)
    assert output.read_bytes() == b"concurrent writer"
    assert not list(workspace.glob("*.partial"))


def test_publish_failure_cleans_staging(workspace, monkeypatch):
    def no_link(*_args):
        raise OSError("synthetic unsupported filesystem")

    monkeypatch.setattr(csv_report.os, "link", no_link)
    output = workspace / "report.xlsx"
    with pytest.raises(OSError):
        csv_report._publish(BytesIO(b"complete report"), output)
    assert not output.exists()
    assert not list(workspace.glob("*.partial"))


def test_default_output_is_unique_localappdata_and_as_of_defaults_to_utc(workspace, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(workspace))
    source = export(workspace, [record(**{KEY + "rotatedAt": ""})])
    before = datetime.now(UTC).date()
    first = generate_report([source])
    second = generate_report([source])
    after = datetime.now(UTC).date()
    assert first.path != second.path
    assert first.path.parent == workspace / "CloudSecurityRemediations" / "reports"
    assert any(day.isoformat() in first.path.name for day in (before, after))
    monkeypatch.delenv("LOCALAPPDATA")
    with pytest.raises(ValueError, match="LOCALAPPDATA"):
        generate_report([source])
    monkeypatch.setenv("LOCALAPPDATA", str(REPOSITORY_ROOT))
    with pytest.raises(ValueError, match="outside"):
        generate_report([source])


@pytest.mark.parametrize("as_of", ["20260401", "2026-02-30", "2026-04-01T00:00:00Z"])
def test_bad_as_of(workspace, as_of):
    with pytest.raises(ValueError, match="As-of"):
        generate_report(["unused.csv"], output=workspace / "report.xlsx", as_of=as_of)


def test_invalid_api_context_and_missing_paths(workspace):
    with pytest.raises(ValueError, match="Severity"):
        generate_report(["unused.csv"], output=workspace / "report.xlsx", severity="invented")
    with pytest.raises(ValueError, match="one or more"):
        generate_report([], output=workspace / "report.xlsx")
    with pytest.raises(ValueError, match="sequence"):
        generate_report("unused.csv", output=workspace / "report.xlsx")
    assert csv_report._as_of(date(2026, 4, 1)) == date(2026, 4, 1)


def test_row_and_column_bounds_without_large_allocations(workspace, monkeypatch):
    source = export(workspace, [record(), record()])
    monkeypatch.setattr(csv_report, "_ROW_LIMIT", 1)
    with pytest.raises(ValueError, match="row limit"):
        generate_report([source], output=workspace / "report.xlsx", as_of="2026-04-01")
    monkeypatch.setattr(csv_report, "_ROW_LIMIT", 100)
    monkeypatch.setattr(csv_report, "_COLUMN_LIMIT", 4)
    with pytest.raises(ValueError, match="column limit"):
        generate_report([source], output=workspace / "report.xlsx", as_of="2026-04-01")


def test_long_header_cannot_be_silently_truncated(workspace):
    source = export(workspace, [record(**{"x" * 32768: "value"})])
    with pytest.raises(ValueError, match="header exceeds"):
        generate_report([source], output=workspace / "report.xlsx")


@pytest.mark.parametrize("module", ["remediation", "remediation.csv_report"])
def test_successful_module_cli_prints_only_path_and_counts(workspace, module):
    source = export(workspace, [record()])
    notes = workspace / "notes.txt"
    notes.write_text("Synthetic private analysis", encoding="utf-8")
    output = workspace / "report.xlsx"
    result = subprocess.run(
        [
            sys.executable, "-m", module, str(source),
            "--output", str(output), "--as-of", "2026-04-01",
            "--severity", "Medium", "--rule-id", "Synthetic rule", "--notes", str(notes),
        ],
        cwd=REPOSITORY_ROOT, capture_output=True, text=True, check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    assert str(output) in result.stdout
    assert "Source rows: 1" in result.stdout
    assert "unique credentials: 1" in result.stdout
    assert "Synthetic" not in result.stdout
    assert "graph-key-one" not in result.stdout
    workbook = load_workbook(output)
    cover_text = "\n".join(str(cell.value) for row in workbook["Cover"] for cell in row)
    assert "Medium" in cover_text
    assert "Synthetic rule" in cover_text
    assert not workbook["Cover"]._charts
    for name in ("AWS Data", "Azure Data"):
        headers = [cell.value for cell in workbook[name][1]]
        assert not {"Severity", "Rule ID", "Wiz ID", "Issue ID"}.intersection(headers)
    assert "Synthetic private analysis" in workbook["Remediation"]["A1"].comment.text


def test_reference_style_cover_and_nine_column_resource_tables(workspace):
    _, workbook = report(workspace, [record()])
    cover = workbook["Cover"]
    assert not cover.merged_cells
    assert cover["A1"].fill.fgColor.rgb.endswith("1A3D5C")
    assert cover["A1"].font.sz == 28
    assert cover["A5"].value == "Executive Summary"
    assert cover["A11"].value == "Issue Summary"
    assert cover["A16"].value == "Risk Description"
    assert cover["A19"].value == "Security Recommendation and Requested Response"
    assert not cover.sheet_view.showGridLines
    assert cover.page_setup.fitToHeight == 1
    for name in ("AWS Data", "Azure Data"):
        sheet = workbook[name]
        assert sheet.max_column == 9
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref.endswith(f"I{sheet.max_row}")
        assert sheet["A1"].fill.fgColor.rgb.endswith("1A3D5C")
        assert sheet.page_margins.bottom >= sheet.page_margins.footer + 0.3
        assert sheet.print_title_rows == "$1:$1"
        for row in sheet.iter_rows(min_row=2):
            assert 40 <= sheet.row_dimensions[row[0].row].height <= 409
            assert all(cell.alignment.wrap_text for cell in row)
    remediation = workbook["Remediation"]
    text = "\n".join(str(cell.value) for row in remediation for cell in row if cell.value)
    assert "Approval Required Before Changes" in text
    assert "two-access-key limit" in text
    assert "No Longer Needed" in text
    assert "AWS" in text and "Azure / Entra" in text
    assert "fresh Wiz assessment" in text


def test_summary_keeps_full_values_without_removed_tab_references(workspace):
    name = "Synthetic long identity " * 12
    permission = "Synthetic long permission " * 12
    _, workbook = report(workspace, [
        record(**{PRINCIPAL + "Name": name, PERMISSION + "Name": permission}),
    ])
    summary = sheet_rows(workbook["AWS Data"])[0]
    assert summary["Resource Name"] == name
    assert summary["Permission Summary"] == permission
    assert workbook.sheetnames == csv_report.SHEETS


@pytest.mark.parametrize(
    ("rotation", "expiry"),
    [
        ("2025-03-31", "2026-04-01"),
        ("2025-04-01", "2026-05-31"),
        ("2025-04-02", "2026-06-01"),
        ("2026-03-01", "2026-03-31"),
    ],
)
def test_dates_preserve_accepted_stripe_design(workspace, rotation, expiry):
    _, workbook = report(workspace, [
        record(**{KEY + "rotatedAt": rotation, KEY + "validBefore": expiry}),
    ])
    assert copy(workbook["AWS Data"]["E2"].fill) == copy(workbook["AWS Data"]["A2"].fill)
    assert copy(workbook["AWS Data"]["F2"].fill) == copy(workbook["AWS Data"]["A2"].fill)


def test_cover_counts_unique_keys_instead_of_repeated_scopes(workspace):
    _, workbook = report(workspace, [
        record(**{KEY + "rotatedAt": "2020-01-01"}),
        record(**{KEY + "rotatedAt": "2020-01-01", SCOPE + "ID": "second-scope"}),
    ])
    assert "1 identities | 1 credentials" in workbook["Cover"]["A7"].value
    assert workbook["Cover"]["A13"].value.startswith("1 credentials")
    assert workbook["Cover"]["A12"].value.startswith("1 credentials")
    assert workbook["AWS Data"].max_row == 3


def issue_record(number=1, **overrides):
    return {
        "Issue ID": f"issue-{number}",
        "Control ID": "synthetic-control",
        "Resource vertex ID": "principal-one",
        "Resource external ID": "native-principal",
        "Resource Platform": "AWS",
        "Title": "Synthetic credential finding",
        "Severity": "MEDIUM",
        "Status": "IN_PROGRESS",
        "Wiz URL": f"https://app.wiz.io/issues#~(issue~'issue-{number})",
        "Updated At": "2026-09-01T00:00:00Z",
        **overrides,
    }


def test_issue_enrichment_is_wired_to_report_and_never_approval(workspace):
    source = export(workspace, [
        record(), record(**{KEY + "ID": "graph-key-two"}),
        record(**{SCOPE + "ID": "second-scope", SCOPE + "externalId": "scope-alternative"}),
    ])
    issue = issue_record()
    issue_path = export(workspace, [issue], "issues.csv")
    result = generate_report(
        [source], issues=[issue_path], output=workspace / "report.xlsx", as_of="2026-04-01",
    )
    assert (result.unique_issues, result.unique_principals, result.unique_credentials) == (1, 1, 2)
    assert result.grouped_scopes == 3
    workbook = load_workbook(result.path)
    assert workbook.sheetnames == csv_report.SHEETS
    assert workbook["Cover"]["A6"].value == issue["Title"]
    assert "1 identities | 2 credentials" in workbook["Cover"]["A7"].value
    assert "1 Wiz issue" in workbook["Cover"]["A9"].value
    assert "1 Wiz issues" not in workbook["Cover"]["A9"].value
    assert "Medium" in workbook["Cover"]["A23"].value
    sheet = workbook["AWS Data"]
    assert sheet["I1"].value == "Wiz Status / Approval"
    for row in range(2, 5):
        assert sheet.cell(row, 2).hyperlink.target == issue["Wiz URL"]
        assert "IN_PROGRESS" in sheet.cell(row, 2).comment.text
        assert issue["Updated At"] in sheet.cell(row, 2).comment.text
        assert "Approval not recorded" in sheet.cell(row, 9).value
        assert sheet.cell(row, 5).value == "90"


@pytest.mark.parametrize("override", [
    {"title": "Contradictory title"}, {"severity": "High"}, {"rule_id": "different-control"},
])
def test_conflicting_issue_context_never_publishes(workspace, override):
    graph = export(workspace, [record()])
    issue = export(workspace, [issue_record()], "issues.csv")
    with pytest.raises(ValueError, match="conflicts"):
        generate_report(
            [graph], issues=[issue], output=workspace / "report.xlsx",
            as_of="2026-04-01", **override,
        )
    assert not (workspace / "report.xlsx").exists()


def test_blank_issue_title_uses_default_or_provided_context(workspace):
    graph = export(workspace, [record()])
    issue = export(workspace, [issue_record(Title="")], "issues.csv")
    result = generate_report(
        [graph], issues=[issue], output=workspace / "report.xlsx",
        as_of="2026-04-01", title="Provided finding context",
    )
    assert load_workbook(result.path)["Cover"]["A6"].value == "Provided finding context"


@pytest.mark.parametrize("value", ["issues.csv", Path("issues.csv"), []])
def test_issue_api_requires_nonempty_sequence(workspace, value):
    graph = export(workspace, [record()])
    with pytest.raises(ValueError, match="issue CSV paths"):
        generate_report([graph], issues=value, output=workspace / "report.xlsx")


def test_explicit_update_reuses_path_and_rebuilds_only_four_tabs(workspace):
    result, workbook = report(workspace, [record()])
    workbook["AWS Data"]["I2"] = "Manually edited decision - not a source of authorization"
    workbook.save(result.path)
    previous = result.path.read_bytes()
    graph = workspace / "input.csv"
    issue = export(workspace, [issue_record()], "issues.csv")
    updated = generate_report(
        [graph], issues=[issue], output=result.path, as_of="2026-04-02", update=True,
    )
    assert updated.path == result.path
    assert result.path.read_bytes() != previous
    assert sorted(path.name for path in workspace.glob("*.xlsx")) == ["report.xlsx"]
    checked = load_workbook(updated.path)
    assert checked.sheetnames == csv_report.SHEETS
    assert checked["AWS Data"]["E2"].value == "91"
    assert checked["AWS Data"]["I2"].value == "In Progress\nApproval not recorded"
    assert not list(workspace.glob("*.partial"))


@pytest.mark.parametrize("failure", ["parse", "metadata", "layout", "replace"])
def test_update_failure_preserves_original_file(workspace, monkeypatch, failure):
    result, _ = report(workspace, [record()])
    original = result.path.read_bytes()
    graph = workspace / "input.csv"
    kwargs = {}
    if failure == "parse":
        graph.write_text("broken", encoding="utf-8")
    elif failure == "metadata":
        issue = export(
            workspace, [issue_record(**{"Resource external ID": "mismatch"})], "issues.csv"
        )
        kwargs["issues"] = [issue]
    elif failure == "layout":
        export(workspace, [record(**{PRINCIPAL + "Name": "x" * 32768})])
    else:
        def locked(*_args):
            raise PermissionError("synthetic locked workbook")
        monkeypatch.setattr(csv_report.os, "replace", locked)
    with pytest.raises((ValueError, PermissionError)):
        generate_report([graph], output=result.path, as_of="2026-04-01", update=True, **kwargs)
    assert result.path.read_bytes() == original
    assert not list(workspace.glob("*.partial"))


def test_update_requires_explicit_existing_tool_workbook(workspace):
    graph = export(workspace, [record()])
    for output, match in ((None, "explicit"), (workspace / "missing.xlsx", "existing")):
        with pytest.raises(ValueError, match=match):
            generate_report([graph], output=output, update=True)
    foreign = workspace / "foreign.xlsx"
    Workbook().save(foreign)
    original = foreign.read_bytes()
    with pytest.raises(ValueError, match="from this tool"):
        generate_report([graph], output=foreign, update=True)
    assert foreign.read_bytes() == original
    foreign.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="not a valid"):
        generate_report([graph], output=foreign, update=True)


def test_update_detects_change_during_generation(workspace, monkeypatch):
    result, _ = report(workspace, [record()])
    original_workbook = csv_report._workbook

    def concurrent_edit(*args, **kwargs):
        stream = original_workbook(*args, **kwargs)
        result.path.write_bytes(b"concurrent writer")
        return stream

    monkeypatch.setattr(csv_report, "_workbook", concurrent_edit)
    with pytest.raises(ValueError, match="changed during"):
        generate_report(
            [workspace / "input.csv"], output=result.path, update=True, as_of="2026-04-01"
        )
    assert result.path.read_bytes() == b"concurrent writer"
    assert not list(workspace.glob("*.partial"))


@pytest.mark.parametrize("malformed", [False, True])
def test_update_rejects_entity_expansion_and_malformed_xml(workspace, malformed):
    result, _ = report(workspace, [record()])
    rebuilt = BytesIO()
    with ZipFile(result.path) as original, ZipFile(rebuilt, "w") as changed:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = b"<broken" if malformed else (
                    b'<!DOCTYPE coreProperties [<!ENTITY synthetic "expanded">]>'
                    b'<coreProperties xmlns="http://schemas.openxmlformats.org/'
                    b'package/2006/metadata/core-properties">&synthetic;</coreProperties>'
                )
            changed.writestr(item, data)
    original_bytes = rebuilt.getvalue()
    result.path.write_bytes(original_bytes)
    with pytest.raises(ValueError, match="not a valid"):
        generate_report([workspace / "input.csv"], output=result.path, update=True)
    assert result.path.read_bytes() == original_bytes


def test_cli_repeated_issue_options_and_update_warning(workspace, capsys):
    result, _ = report(workspace, [record()])
    issue = export(workspace, [issue_record()], "issues.csv")
    status = main([
        str(workspace / "input.csv"), "--issues", str(issue), "--issues", str(issue),
        "--output", str(result.path), "--update", "--as-of", "2026-04-01",
    ])
    captured = capsys.readouterr()
    assert status == 0
    assert "Unique Wiz issues: 1" in captured.out
    assert "manual workbook edits will be replaced" in captured.err
    assert "Synthetic" not in captured.out


@pytest.mark.parametrize("field", [KEY + "active", PRINCIPAL + "hasHighPrivileges"])
def test_cross_scope_entity_drift_is_rejected(workspace, field):
    graph = export(workspace, [
        record(),
        record(**{SCOPE + "ID": "different-scope", field: "false"}),
    ])
    with pytest.raises(ValueError, match="Conflicting entity evidence"):
        generate_report([graph], output=workspace / "report.xlsx", as_of="2026-04-01")
