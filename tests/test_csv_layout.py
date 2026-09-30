"""Synthetic-only checks for the reusable four-tab spreadsheet layout."""

from collections.abc import Sequence
from datetime import date
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from remediation.csv_layout import (
    MAX_ROWS,
    MAX_TEXT,
    TABLE_FIELDS,
    _height,
    cover,
    remediation_page,
    resource_table,
)


def record(**overrides):
    result = {
        "cloud": "AWS",
        "scope_name": "Synthetic scope",
        "scope_external_id": "scope-native-id",
        "scope_id": "scope-wiz-id",
        "scope_type": "subscription",
        "principal": "Synthetic workload",
        "principal_native_id": "principal-native-id",
        "principal_id": "principal-wiz-id",
        "principal_type": "synthetic-type",
        "credential_name": "Synthetic credential",
        "credential_id": "credential-id",
        "key_active": "true",
        "high_privileges": "true",
        "inactive": "Unknown",
        "rotated_at": "2024-01-01",
        "age_days": "500",
        "valid_before": "2026-01-15",
        "days_until_expiry": "14",
        "permission_names": "Read synthetic objects\nUpdate synthetic objects",
        "permission_access_types": "read,write",
        "source_rows": "2,3",
        "statusDraft": "Draft",
    }
    result.update(overrides)
    return result


def empty_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def add_cover(workbook, records=(), **kwargs):
    options = {
        "title": "Synthetic privileged credential finding",
        "as_of": date(2026, 1, 1),
        "source_rows": 15,
        "severity": "High",
        "rule_id": "synthetic-rule",
    }
    options.update(kwargs)
    cover(workbook, records, **options)


def assert_font(cell, size, color, *, bold=False):
    assert cell.font.name == "Calibri"
    assert cell.font.sz == size
    assert cell.font.b == bold
    assert cell.font.color.type == "rgb"
    assert cell.font.color.rgb[-6:] == color


def assert_fill(cell, color):
    assert cell.fill.patternType == "solid"
    assert cell.fill.fgColor.rgb[-6:] == color


def assert_border(cell):
    for side in ("left", "right", "top", "bottom"):
        edge = getattr(cell.border, side)
        assert edge.style == "thin"
        assert edge.color.rgb[-6:] == "D2DCE5"


def roundtrip(workbook):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return load_workbook(stream), stream


def test_four_tab_layout_roundtrip_and_single_column_pages():
    workbook = empty_workbook()
    add_cover(workbook, [record()])
    resource_table(workbook, "AWS Data", [record()])
    resource_table(workbook, "Azure Data", [])
    remediation_page(workbook)
    workbook, _ = roundtrip(workbook)
    assert workbook.sheetnames == ["Cover", "AWS Data", "Azure Data", "Remediation"]
    for name in ("Cover", "Remediation"):
        sheet = workbook[name]
        assert not sheet.merged_cells.ranges
        assert sheet.max_column == 1
        assert sheet.column_dimensions["A"].width == 100
        assert sheet.sheet_view.zoomScale == 100
        assert sheet.page_setup.orientation == "portrait"
        assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
        for address, size, color in (("A1", 28, "1A3D5C"), ("A3", 16, "2E5266")):
            cell = sheet[address]
            assert_font(cell, size, "FFFFFF", bold=True)
            assert_fill(cell, color)
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
            assert cell.alignment.wrap_text
        assert sheet.row_dimensions[1].height >= 70.05
        assert sheet.row_dimensions[3].height >= 30
    assert workbook["Cover"].page_setup.fitToHeight == 1
    assert workbook["Remediation"].page_setup.fitToHeight == 0
    assert [item.id for item in workbook["Remediation"].row_breaks.brk] == [19]
    assert workbook["Remediation"]["A17"].value.startswith("Step 4:")
    assert workbook["Remediation"]["A20"].value.startswith("Step 5:")


def test_cover_exact_approved_colors_fonts_borders_and_rows():
    workbook = empty_workbook()
    add_cover(workbook, [record()])
    sheet = workbook["Cover"]
    assert sheet["A5"].value == "Executive Summary"
    assert_font(sheet["A5"], 13, "FFFFFF", bold=True)
    assert_fill(sheet["A5"], "003D7A")
    for row in (6, 7, 8, 9, 10, 23):
        cell = sheet.cell(row, 1)
        assert_font(cell, 11, "003D7A", bold=True)
        assert_fill(cell, "E8EFF7")
        assert_border(cell)
        assert cell.alignment.wrap_text
    for row, label, color in (
        (11, "Issue Summary", "CC3333"),
        (16, "Risk Description", "FF9900"),
        (19, "Security Recommendation and Requested Response", "00A651"),
    ):
        assert sheet.cell(row, 1).value == label
        assert_font(sheet.cell(row, 1), 12, "FFFFFF", bold=True)
        assert_fill(sheet.cell(row, 1), color)
    for row, color in ((12, "FFE8E8"), (13, "FFE8E8"), (14, "FFE8E8"),
                       (17, "FFF4E6"), (20, "E8F5E8")):
        cell = sheet.cell(row, 1)
        assert_font(cell, 10, "000000")
        assert_fill(cell, color)
        assert cell.alignment.horizontal == "left"
        assert cell.alignment.vertical == "top"
        assert cell.alignment.wrap_text
    assert "2026-01-01 (UTC)" in sheet["A10"].value


def test_table_exact_styles_stripes_printing_and_no_stale_panes():
    workbook = empty_workbook()
    resource_table(workbook, "AWS Data", [
        record(credential_id="credential-one"),
        record(credential_id="credential-two", days_until_expiry="-2"),
    ])
    workbook, stream = roundtrip(workbook)
    sheet = workbook["AWS Data"]
    assert sheet.max_column == 9
    assert sheet.freeze_panes == "A2"
    assert sheet.sheet_view.zoomScale == 100
    selections = sheet.sheet_view.selection
    assert len(selections) == 1
    assert (selections[0].pane, selections[0].activeCell, selections[0].sqref) == (
        "bottomLeft", "A2", "A2",
    )
    assert sheet.sheet_view.pane.xSplit is None
    assert sheet.sheet_view.pane.ySplit == 1
    assert sheet.sheet_view.pane.activePane == "bottomLeft"
    assert sheet.page_setup.orientation == "landscape"
    assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A3
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.auto_filter.ref == "A1:I3"
    assert "$A$1:$I$3" in str(sheet.print_area)
    for column, (_, label, width) in enumerate(TABLE_FIELDS, 1):
        cell = sheet.cell(1, column)
        assert cell.value == label
        assert sheet.column_dimensions[cell.column_letter].width == width
        assert_font(cell, 11, "FFFFFF", bold=True)
        assert_fill(cell, "1A3D5C")
        assert_border(cell)
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "center"
        assert cell.alignment.wrap_text
    for row, color in ((2, "D9E8F5"), (3, "FFFFFF")):
        for cell in sheet[row]:
            assert_font(cell, 10, "000000")
            assert_fill(cell, color)
            assert_border(cell)
            assert cell.alignment.horizontal == "left"
            assert cell.alignment.vertical == "center"
            assert cell.alignment.wrap_text
    with ZipFile(stream) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert xml.count(b"<selection ") == 1
        assert b"bottomRight" not in xml
        assert b"topRight" not in xml


def test_full_names_permissions_and_source_context_are_preserved():
    name = " ".join(["CompleteSyntheticName"] * 15)
    permissions = "\n".join(f"FullSyntheticPermission{index}" for index in range(15))
    workbook = empty_workbook()
    resource_table(workbook, "AWS Data", [
        record(principal=name, credential_name=name, permission_names=permissions),
    ])
    sheet = workbook["AWS Data"]
    assert sheet["B2"].value == name
    assert sheet["C2"].value == name
    assert sheet["G2"].value == permissions
    assert sheet.row_dimensions[2].height > 200
    assert sheet.row_dimensions[2].height <= 409
    context = sheet["B2"].comment.text
    assert sheet["A2"].value == "Synthetic scope\nscope-native-id"
    for value in ("scope-native-id", "scope-wiz-id", "principal-native-id", "principal-wiz-id",
                  "credential-id", "2,3", "2024-01-01", "read,write"):
        assert value in context
    assert "Source CSV files remain the original evidence." in context
    assert "Credential Details" not in context
    assert sheet["B2"].hyperlink is None


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_cells_and_notes_are_literal_after_serialization(prefix):
    text = prefix + "SYNTHETIC(1)"
    workbook = empty_workbook()
    add_cover(workbook, [record()], title=text, rule_id=text)
    resource_table(workbook, "AWS Data", [
        record(principal=text, credential_name=text, permission_names=text),
    ])
    remediation_page(workbook, notes=text)
    workbook, _ = roundtrip(workbook)
    for name, addresses in (("Cover", ("A6",)), ("AWS Data", ("B2", "C2", "G2"))):
        for address in addresses:
            assert workbook[name][address].value == text
            assert workbook[name][address].data_type == "s"
    assert text in workbook["Remediation"]["A1"].comment.text
    assert "not validated and not approval" in workbook["Remediation"]["A1"].comment.text
    assert text in workbook["Cover"]["A23"].value


def test_missing_fields_unknown_severity_and_no_invented_issue_data():
    workbook = empty_workbook()
    unknown = {"cloud": "Azure", "principal_id": "synthetic-principal", "credential_id": "key"}
    add_cover(workbook, [unknown], severity="not-a-source-severity")
    resource_table(workbook, "Azure Data", [unknown])
    sheet = workbook["Azure Data"]
    assert sheet["A2"].value == "Unknown"
    assert sheet["B2"].value == "Unknown"
    assert sheet["D2"].value == "Unknown"
    assert sheet["F2"].value == "Not supplied"
    assert sheet["H2"].value == "Unknown"
    assert sheet["I1"].value == "Remediation Status"
    assert sheet["I2"].value == "Pending review\nApproval not recorded"
    assert sheet["B2"].hyperlink is None
    assert "Issue ID:" not in sheet["B2"].comment.text
    assert "Source issue severity: Unknown" in workbook["Cover"]["A23"].value
    assert "Provided context (not verified): severity Unknown" in workbook["Cover"]["A23"].value
    assert "Wiz issue count unknown" in workbook["Cover"]["A9"].value
    assert workbook["Cover"]["A12"].value.startswith("0 credentials")


def test_unique_issue_counts_severity_and_lifecycle_are_separate_from_credentials():
    issues = [
        {"issue_id": f"issue-{index}", "issue_severity": "Low",
         "issue_status": "Open" if index < 4 else "In Progress",
         "issue_control_id": "verified-control"}
        for index in range(8)
    ]
    records = [
        record(credential_id=f"key-{index}", **issue) for index, issue in enumerate(issues)
    ]
    records.append(record(credential_id="second-key", **issues[0]))
    workbook = empty_workbook()
    add_cover(workbook, records, severity="Critical", issues=issues)
    sheet = workbook["Cover"]
    assert "9 credentials" in sheet["A9"].value
    assert "8 Wiz issues" in sheet["A9"].value
    assert "15 source CSV rows" in sheet["A9"].comment.text
    assert "9 credential-scope rows" in sheet["A9"].comment.text
    assert "Wiz severity: Low: 8" in sheet["A23"].value
    assert "Critical" not in sheet["A23"].value
    context = sheet["A23"].comment.text
    assert "derived from matched issue CSV records" in context
    assert "example" not in context.casefold()
    assert "Control ID: verified-control" in sheet["A23"].value
    assert "In Progress: 4, Open: 4" in sheet["A23"].value
    assert "In Progress: 4, Open: 4" in context
    derived = empty_workbook()
    add_cover(derived, records)
    assert "8 Wiz issues" in derived["Cover"]["A9"].value
    assert "Low: 8" in derived["Cover"]["A23"].value


@pytest.mark.parametrize("issues", [None, [], [
    {"issue_id": "synthetic-issue", "issue_severity": "Medium", "issue_status": "Open"},
]])
def test_stakeholder_summary_keeps_collection_details_in_notes(issues):
    workbook = empty_workbook()
    add_cover(workbook, [record()], issues=issues)
    workbook, _ = roundtrip(workbook)
    sheet = workbook["Cover"]
    visible = "\n".join(str(cell.value) for row in sheet for cell in row if cell.value)
    for technical_phrase in ("source CSV rows", "credential-scope rows", "matched issue CSV"):
        assert technical_phrase not in visible
    assert "15 source CSV rows; 1 credential-scope rows" in sheet["A9"].comment.text
    assert "current configuration still needs confirmation" in sheet["A10"].comment.text
    assert "current configuration still needs confirmation" not in visible
    assert "Remediation approval: Pending" in visible
    assert "We recommend staged replacement" in visible
    assert "Security recommends" not in visible
    assert "Please confirm the preferred approach and any business constraints" in visible
    assert "Please confirm ownership, required access and application dependencies" in visible
    assert "An inactivity flag alone is not a reason to remove a credential" in (
        sheet["A14"].comment.text
    )
    assert "inactivity flag alone" not in visible
    assert "Missing information is marked" not in visible
    assert "These credentials authenticate AWS or Azure identities" in sheet["A14"].value
    assert "No approval, business risk acceptance, execution or verified closure" in (
        sheet["A23"].comment.text
    )
    if issues is None:
        assert "Provided context (not verified)" in visible


@pytest.mark.parametrize(("workflow_text", "workflow_context"), [
    ("Stage 1/8 | Evidence review in progress", ""),
    ("", "Owner enrichment incomplete; reconcile source rows"),
    ("Status: blocked", "Confirm ticket routing; fetch graph"),
])
def test_internal_workflow_is_note_only_and_preserves_stakeholder_request(
    workflow_text, workflow_context,
):
    workbook = empty_workbook()
    add_cover(workbook, [record()], workflow_text=workflow_text, workflow_context=workflow_context)
    workbook, _ = roundtrip(workbook)
    sheet = workbook["Cover"]
    visible = "\n".join(str(c.value) for row in sheet for c in row if c.value)
    for text in (workflow_text, workflow_context):
        if text:
            assert text not in visible
            assert text in sheet["A21"].comment.text
    assert "Please confirm ownership, required access and application dependencies" in (
        sheet["A21"].value
    )
    assert sheet["A16"].value == "Risk Description"
    assert "not evidence of a compromise or data loss" in sheet["A17"].value


def test_aged_findings_are_conditional_and_expiry_is_not_an_sla():
    records = [
        record(credential_id="aged"),
        record(credential_id="inactive", key_active="false"),
        record(credential_id="unprivileged", high_privileges="false"),
        record(credential_id="boundary", age_days="365"),
        record(credential_id="unknown", age_days="Unknown"),
        record(credential_id="expired", age_days="100", days_until_expiry="-1"),
    ]
    workbook = empty_workbook()
    add_cover(workbook, records)
    sheet = workbook["Cover"]
    assert sheet["A12"].value.startswith("1 credentials")
    assert "1 have reported expiry dates in the past" in sheet["A13"].value
    assert "2026-01-15" in sheet["A13"].value
    assert "Expiry dates are not remediation deadlines" in sheet["A13"].comment.text
    assert "remediation deadlines" not in sheet["A13"].value
    assert "not evidence of a compromise or data loss" in sheet["A17"].value
    assert "do not prove unrestricted cloud access" in sheet["A17"].value


@pytest.mark.parametrize("status", ["Open", "In Progress", "Resolved", "Unknown"])
def test_enrichment_link_preserves_body_font_and_lifecycle_is_not_approval(status):
    workbook = empty_workbook()
    resource_table(workbook, "AWS Data", [record(
        issue_id="synthetic-issue",
        issue_status=status,
        issue_status_raw=status.upper().replace(" ", "_"),
        issue_url="https://app.wiz.io/issues#~(issue~'synthetic-issue)",
        issue_updated_at="2026-01-01T10:00:00Z",
        issue_due_at="2026-03-01T00:00:00Z",
        issue_recommendation="Source-provided proposal",
    )])
    workbook, _ = roundtrip(workbook)
    sheet = workbook["AWS Data"]
    assert sheet["I1"].value == "Wiz Status / Approval"
    assert sheet["I2"].value == status + "\nApproval not recorded"
    assert sheet["B2"].hyperlink.target == "https://app.wiz.io/issues#~(issue~'synthetic-issue)"
    assert_font(sheet["B2"], 10, "000000")
    assert sheet["B2"].font.underline is None
    context = sheet["B2"].comment.text
    assert "not a fresh credential observation" in context
    assert "not approval, risk acceptance or verified closure" in context
    assert "2026-03-01T00:00:00Z" in context


@pytest.mark.parametrize("url", [
    "http://app.wiz.io/issues#~(issue~'synthetic-issue)",
    "https://app.wiz.io:443/issues#~(issue~'synthetic-issue)",
    "https://name:secret" + "@app.wiz.io/issues#~(issue~'synthetic-issue)",
    "https://app.wiz.io.invalid/issues#~(issue~'synthetic-issue)",
    "https://app.wiz.io/issues#~(issue~'different-issue)",
    "https://app.wiz.io/other#~(issue~'synthetic-issue)",
    "https://app.wiz.io/issues?issue=synthetic-issue",
    "https://app.wiz.io/issues#~(issue~'synthetic-issue)\n",
    "https://app.wiz.io/issues#~(issue~'synthetic-issue-extra)",
    "https://app.wiz.io/issues#~(issue~'prefix-synthetic-issue)",
    "=SYNTHETIC(1)",
])
def test_invalid_issue_targets_fail_closed(url):
    workbook = empty_workbook()
    with pytest.raises(ValueError, match="Wiz URL"):
        resource_table(workbook, "AWS Data", [record(
            issue_id="synthetic-issue", issue_url=url,
        )])


def test_nonblank_issue_url_without_source_issue_id_is_rejected():
    workbook = empty_workbook()
    with pytest.raises(ValueError, match="requires a source issue ID"):
        resource_table(workbook, "AWS Data", [record(
            issue_url="https://app.wiz.io/issues#~(issue~'synthetic-issue)",
        )])


@pytest.mark.parametrize("url", [
    "https://app.wiz.io/issues#~(issue~'synthetic-issue)",
    "https://app.wiz.io/issues#%7E%28issue%7E%27synthetic-issue%29",
    "https://app.wiz.io/issues?view=all#~(issue~'synthetic-issue)",
    "https://app.wiz.io/issues/synthetic-issue",
])
def test_wiz_fragment_link_contract_and_decoded_exact_id_tokens(url):
    workbook = empty_workbook()
    resource_table(workbook, "AWS Data", [record(
        issue_id="synthetic-issue", issue_url=url,
    )])
    assert workbook["AWS Data"]["B2"].hyperlink.target == url


def test_cross_scope_rows_do_not_inflate_identity_or_credential_counts():
    records = [
        record(scope_external_id="scope-one"),
        record(scope_external_id="scope-two"),
        record(scope_external_id="scope-three", principal_id="other-principal"),
    ]
    workbook = empty_workbook()
    add_cover(workbook, records)
    assert workbook["Cover"]["A7"].value == "AWS: 2 identities | 1 credentials"
    assert "Total: 2 identities | 1 credentials" in workbook["Cover"]["A9"].value
    assert "3 credential-scope rows" in workbook["Cover"]["A9"].comment.text
    assert workbook["Cover"]["A12"].value.startswith("1 credentials")


@pytest.mark.parametrize("key", ["cloud", "principal_id", "credential_id"])
@pytest.mark.parametrize("renderer", ["cover", "table"])
def test_missing_required_count_identifiers_fail_instead_of_using_names(key, renderer):
    incomplete = record()
    del incomplete[key]
    with pytest.raises(ValueError, match=f"nonblank {key} identifier"):
        if renderer == "cover":
            add_cover(empty_workbook(), [incomplete])
        else:
            resource_table(empty_workbook(), "AWS Data", [incomplete])


def test_graph_and_issue_scope_mismatch_context_is_retained():
    workbook = empty_workbook()
    note = "Issue and graph scopes differ; both source values are preserved."
    resource_table(workbook, "Azure Data", [record(
        issue_scope_note=note, issue_subscription_id="other-scope-native-id",
    )])
    cell = workbook["Azure Data"]["B2"]
    assert note in cell.comment.text
    assert "Scope Wiz ID: scope-wiz-id" in cell.comment.text
    assert "Issue subscription ID: other-scope-native-id" in cell.comment.text


def test_without_issue_csv_provided_control_and_severity_are_visible_unverified_context():
    workbook = empty_workbook()
    add_cover(workbook, [record()], severity="High", rule_id="provided-control")
    footer = workbook["Cover"]["A23"]
    assert "Provided context (not verified): severity High" in footer.value
    assert "Control ID: provided-control" in footer.value
    assert "Source issue severity: Unknown" in footer.value
    assert workbook["Cover"].row_dimensions[23].height <= 409


@pytest.mark.parametrize("character", ["\x00", "\x0b", "\ud800", "\udfff", "\ufffe", "\uffff"])
@pytest.mark.parametrize("target", ["cell", "comment", "notes"])
def test_xml_illegal_characters_fail_before_serialization(character, target):
    text = "Synthetic" + character + "Text"
    workbook = empty_workbook()
    with pytest.raises(ValueError, match="characters not supported by XLSX") as error:
        if target == "notes":
            remediation_page(workbook, notes=text)
        else:
            field = "principal" if target == "cell" else "issue_scope_note"
            resource_table(workbook, "AWS Data", [record(**{field: text})])
    assert "Synthetic" not in str(error.value)


@pytest.mark.parametrize("field", [
    "scope_name", "principal", "credential_name", "permission_names",
    "principal_native_id", "issue_description",
])
def test_overlong_display_and_comment_fields_fail_closed_without_source_disclosure(field):
    workbook = empty_workbook()
    source_marker = "SyntheticSensitiveValue"
    with pytest.raises(ValueError, match="32767-character limit") as error:
        resource_table(workbook, "AWS Data", [
            record(**{field: source_marker + "x" * MAX_TEXT}),
        ])
    assert source_marker not in str(error.value)


@pytest.mark.parametrize("field", ["principal", "credential_name", "permission_names"])
def test_overheight_fields_fail_instead_of_truncating_or_capping(field):
    workbook = empty_workbook()
    with pytest.raises(ValueError, match="409-point row height limit"):
        resource_table(workbook, "AWS Data", [record(**{field: "Fullword " * 150})])


def test_long_title_is_preserved_or_explicitly_rejected():
    workbook = empty_workbook()
    title = "FullSyntheticTitle " * 30
    add_cover(workbook, title=title)
    assert workbook["Cover"]["A6"].value == title
    with pytest.raises(ValueError, match="409-point row height limit"):
        add_cover(empty_workbook(), title="FullSyntheticTitle " * 200)
    with pytest.raises(ValueError, match="32767-character limit"):
        add_cover(empty_workbook(), title="x" * (MAX_TEXT + 1))


def test_notes_are_retained_as_review_only_comments_and_limit_errors_are_explicit():
    workbook = empty_workbook()
    notes = "=Literal synthetic note\nSecond line - no approval recorded."
    remediation_page(workbook, notes=notes)
    sheet = workbook["Remediation"]
    assert sheet["A1"].comment.text.endswith(notes)
    assert "Provided review notes - not validated and not approval" in sheet["A1"].comment.text
    assert workbook.sheetnames == ["Remediation"]
    assert_fill(sheet["A5"], "FF9900")
    assert_fill(sheet["A6"], "FFF4E6")
    for row in (8, 11, 14, 17, 20, 23, 26):
        assert_fill(sheet.cell(row, 1), "00A651")
        assert_font(sheet.cell(row, 1), 12, "FFFFFF", bold=True)
        assert_fill(sheet.cell(row + 1, 1), "E8F5E8")
        assert_font(sheet.cell(row + 1, 1), 10, "000000")
    with pytest.raises(ValueError, match="Review notes exceeds the Excel 32767-character limit"):
        remediation_page(empty_workbook(), notes="n" * (MAX_TEXT + 1))


def test_empty_table_bounds_and_height_boundary():
    workbook = empty_workbook()
    resource_table(workbook, "Azure Data", [])
    sheet = workbook["Azure Data"]
    assert sheet.max_row == 1
    assert sheet.auto_filter.ref == "A1:I1"
    assert "$A$1:$I$1" in str(sheet.print_area)
    assert _height("\n".join(["x"] * 28), 100) == 402
    with pytest.raises(ValueError, match="409-point"):
        _height("\n".join(["x"] * 29), 100)


def test_height_estimate_accounts_for_wide_characters_and_font_size():
    assert _height("W" * 90, 35) > _height("i" * 90, 35)
    assert _height("\u754c" * 90, 35) > _height("i" * 90, 35)
    assert _height("Synthetic words " * 5, 100, 28) > _height(
        "Synthetic words " * 5, 100, 10,
    )


def test_excess_rows_rejected_before_iteration_or_sheet_creation():
    class TooManyRows(Sequence):
        def __len__(self):
            return MAX_ROWS

        def __getitem__(self, index):
            raise AssertionError("Oversized input must not be read")

    workbook = empty_workbook()
    with pytest.raises(ValueError, match="1048576-row worksheet limit"):
        resource_table(workbook, "AWS Data", TooManyRows())
    assert workbook.sheetnames == []
