"""Synthetic execution evidence only; no live commands or client workbooks."""

import csv
import json
from copy import copy
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook as read_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from remediation.execution_update import parse_result, update_report

REPOSITORY = Path(__file__).resolve().parents[1]
HEADERS = [
    "Account",
    "Resource Name",
    "Region",
    "Remediation Status",
    "Local User Action",
    "Historical Notes",
]
FINAL = "RESULT=ACCOUNT_REMOVAL_VERIFIED_HOME_PATHS_PRESERVED_SSH_CONFIG_UNCHANGED"
EXCLUSIONS = "\nKeep: service-user. Excluded: root; emergency-user."


def load_workbook(path):
    # Inspect snapshots so Excel readers cannot retain OS handles during replacement tests.
    return read_workbook(BytesIO(Path(path).read_bytes()))


@pytest.fixture
def workspace(tmp_path):
    assert not tmp_path.is_relative_to(REPOSITORY)
    return tmp_path


def success(vm="synthetic-vm", users=("alice",)):
    lines = [f"VM={vm} EXECUTION=SINGLE_RUN"]
    lines.extend(
        f"ACCOUNT={user} UID={2000 + index} GID={2000 + index} "
        f"HOME=/home/{user} SHELL=/bin/bash"
        for index, user in enumerate(users)
    )
    lines.extend([
        "PRECHECKS=PASSED; proceeding automatically with approved removals.",
        "BACKUP_PATH=/var/lib/cloud-security-remediation/account-removal-synthetic",
    ])
    for user in users:
        lines.extend([f"REMOVING={user}", f"REMOVED_AND_VERIFIED={user}"])
    lines.extend([f"REMOVED_THIS_RUN={','.join(users)}", FINAL])
    return "\n".join(lines)


def stopped(verified=("alice",), succeeded=("alice", "bob")):
    lines = success(users=("alice", "bob")).splitlines()
    lines = lines[:5]
    for user in ("alice", "bob"):
        lines.append(f"REMOVING={user}")
        if user in verified:
            lines.append(f"REMOVED_AND_VERIFIED={user}")
    lines.append("STOP=Post-removal verification failed for bob")
    lines.append(f"REMOVAL_COMMANDS_SUCCEEDED={','.join(succeeded) if succeeded else 'none'}")
    lines.append("RESULT=STOPPED_REVIEW_REQUIRED; changes may be partial. Do not assume rollback.")
    return "\n".join(lines)


def absent(vm="synthetic-vm"):
    return "\n".join([
        f"VM={vm} EXECUTION=SINGLE_RUN",
        "ACCOUNT=alice ALREADY_ABSENT (not proof of who removed it)",
        "PRECHECKS=PASSED; proceeding automatically with approved removals.",
        "RESULT=NO_CHANGES_ALL_TARGETS_ALREADY_ABSENT",
    ])


def row(vm="synthetic-vm", account="synthetic-account", region="synthetic-region", **extra):
    result = {
        "Account": account,
        "Resource Name": vm,
        "Region": region,
        "Remediation Status": "Pending remediation",
        "Local User Action": "Remove only: alice" + EXCLUSIONS,
        "Historical Notes": "Preserve synthetic history.",
    }
    result.update(extra)
    return result


def workbook(workspace, rows=None, headers=None, cover=True):
    rows = [row()] if rows is None else rows
    headers = HEADERS if headers is None else headers
    book = Workbook()
    book.remove(book.active)
    if cover:
        book.create_sheet("Cover").append(["Synthetic report", "Do not rewrite"])
    for name in ("AWS Data", "Azure Data", "Pending Review", "Pending Remediation"):
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for values in rows:
            if name == "AWS Data":
                continue
            status = values["Remediation Status"].casefold()
            if name == "Pending Review" and status != "pending review":
                continue
            if name == "Pending Remediation" and status != "pending remediation":
                continue
            sheet.append([values.get(header, "") for header in headers])
    path = workspace / "report.xlsx"
    book.save(path)
    book.close()
    return path


def evidence(workspace, text=None, name="result.txt"):
    path = workspace / name
    path.write_text(success() if text is None else text, encoding="utf-8")
    return path


def manual(workspace, updates, headers=None):
    fields = headers or [
        "Account", "Resource Name", "Region", "Status", "Blocker", "Next Action",
    ]
    path = workspace / "blockers.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(updates)
    return path


def blocker(**changes):
    result = {
        "Account": "synthetic-account",
        "Resource Name": "synthetic-vm",
        "Region": "synthetic-region",
        "Status": "Blocked",
        "Blocker": "Owner validation outstanding",
        "Next Action": "Obtain application owner validation",
    }
    result.update(changes)
    return result


def records(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], values_row, strict=True)) for values_row in values[1:]]


def assert_unchanged(report, original, output):
    assert report.read_bytes() == original
    assert not output.exists()


@pytest.mark.parametrize("collapsed", [False, True])
def test_parser_accepts_complete_success_with_empty_stderr(collapsed):
    text = success()
    text = " ".join(text.splitlines()) if collapsed else text
    parsed = parse_result(f"Enable succeeded:\n[stdout]\n{text}\n[stderr]\n")
    assert parsed.vm == "synthetic-vm"
    assert parsed.outcome == "verified"
    assert parsed.verified == ("alice",)
    assert isinstance(parsed.succeeded, tuple)
    assert parsed.absent == ()
    assert parsed.backup == "/var/lib/cloud-security-remediation/account-removal-synthetic"
    assert isinstance(parsed.reason, str)


def test_parser_stopped_only_explicit_verification_proves_removal():
    parsed = parse_result(stopped())
    assert parsed.outcome == "stopped"
    assert parsed.verified == ("alice",)
    assert parsed.succeeded == ("alice", "bob")
    assert "bob" in parsed.reason


def test_parser_stop_without_verified_marker_does_not_infer_success():
    parsed = parse_result(stopped(verified=(), succeeded=("alice", "bob")))
    assert parsed.outcome == "stopped"
    assert parsed.verified == ()
    assert set(parsed.succeeded) == {"alice", "bob"}


def test_parser_already_absent_is_not_a_new_removal():
    parsed = parse_result(absent())
    assert parsed.outcome == "already_absent"
    assert parsed.absent == ("alice",)
    assert parsed.verified == ()
    assert parsed.succeeded == ()


@pytest.mark.parametrize(
    "text",
    [
        "",
        "VM=synthetic-vm EXECUTION=SINGLE_RUN",
        success().replace("EXECUTION=SINGLE_RUN", "EXECUTION=DRY_RUN"),
        success().replace("ACCOUNT=alice UID=2000 GID=2000 HOME=/home/alice SHELL=/bin/bash\n", ""),
        success().replace(" UID=2000", ""),
        success().replace(" GID=2000", ""),
        success().replace(" HOME=/home/alice", ""),
        success().replace(" SHELL=/bin/bash", ""),
        success().replace(
            "PRECHECKS=PASSED; proceeding automatically with approved removals.\n", ""
        ),
        success().replace(
            "BACKUP_PATH=/var/lib/cloud-security-remediation/account-removal-synthetic\n", ""
        ),
        success().replace("REMOVING=alice\n", ""),
        success().replace("REMOVED_AND_VERIFIED=alice\n", ""),
        success().replace("REMOVED_THIS_RUN=alice\n", ""),
        success().replace(FINAL, ""),
        success() + "\nVM=different-vm EXECUTION=SINGLE_RUN",
        success() + "\nSTOPPED=Conflicting final state",
        success().replace("REMOVED_THIS_RUN=alice", "REMOVED_THIS_RUN=bob"),
        success() + "\n[stderr]\nuserdel: removal failed",
        success() + "\nWARNING: unexpected verification failure",
    ],
)
def test_parser_rejects_malformed_truncated_or_conflicting_evidence(text):
    with pytest.raises(ValueError):
        parse_result(text)


def test_full_success_updates_only_target_and_pending_counts_with_backup(workspace):
    report = workbook(workspace, [
        row(),
        row("untouched-vm"),
        row("review-vm", **{"Remediation Status": "Pending review"}),
    ])
    original = report.read_bytes()
    output = workspace / "email.txt"
    result = update_report(report, results=[evidence(workspace)], email_output=output)
    assert result.imported == 1
    assert result.duplicates == 0
    assert Path(result.backup_path).read_bytes() == original
    book = load_workbook(report)
    target, untouched, review = records(book["Azure Data"])
    assert target["Remediation Status"] == "Remediated"
    assert target["Local User Action"].startswith("Removed: alice (Run Command verified).")
    assert target["Local User Action"].endswith(EXCLUSIONS)
    assert untouched["Remediation Status"] == "Pending remediation"
    assert review["Remediation Status"] == "Pending review"
    assert [r["Resource Name"] for r in records(book["Pending Remediation"])] == ["untouched-vm"]
    assert [r["Resource Name"] for r in records(book["Pending Review"])] == ["review-vm"]
    assert book.sheetnames == [
        "Cover", "AWS Data", "Azure Data", "Pending Review", "Pending Remediation",
    ]
    memory = json.loads(result.memory_path.read_text(encoding="utf-8"))
    assert memory["events"][0]["raw_output"] == success()
    assert "BACKUP_PATH" not in str(list(book["Azure Data"].values))
    assert book["Cover"]["B1"].value == "Do not rewrite"
    email = output.read_text(encoding="utf-8")
    assert all(value in email for value in ("synthetic-vm", "alice", "service-user", "root"))
    assert "closed" not in email.casefold()
    assert "password" not in email.casefold()
    assert all("wiz" not in str(cell.value).casefold() for cell in book["Azure Data"][1])
    book.close()


def test_duplicate_evidence_persists_between_updates_and_is_not_new_proof(workspace):
    report = workbook(workspace)
    source = evidence(workspace)
    first = update_report(report, results=[source], email_output=workspace / "first.txt")
    evidence_count = len(json.loads(first.memory_path.read_text())["events"])
    second = update_report(report, results=[source], email_output=workspace / "second.txt")
    assert (first.imported, second.imported, second.duplicates) == (1, 0, 1)
    book = load_workbook(report)
    assert len(json.loads(second.memory_path.read_text())["events"]) == evidence_count
    assert records(book["Azure Data"])[0]["Remediation Status"] == "Remediated"
    book.close()


def test_header_reorder_and_unrelated_style_comment_link_are_preserved(workspace):
    headers = list(reversed(HEADERS))
    report = workbook(workspace, headers=headers, cover=False)
    book = load_workbook(report)
    cell = book["Azure Data"].cell(2, headers.index("Historical Notes") + 1)
    cell.font = Font(bold=True, color="123456")
    cell.fill = PatternFill("solid", fgColor="ABCDEF")
    cell.comment = Comment("Synthetic historical comment", "Synthetic author")
    cell.hyperlink = "https://example.invalid/history"
    expected_style = copy(cell._style)
    book.save(report)
    book.close()
    update_report(report, results=[evidence(workspace)], email_output=workspace / "email.txt")
    book = load_workbook(report)
    target = records(book["Azure Data"])[0]
    assert target["Remediation Status"] == "Remediated"
    cell = book["Azure Data"].cell(2, headers.index("Historical Notes") + 1)
    assert cell.value == "Preserve synthetic history."
    assert cell._style == expected_style
    assert cell.comment.text == "Synthetic historical comment"
    assert cell.hyperlink.target == "https://example.invalid/history"
    assert "Cover" not in book.sheetnames
    book.close()


def test_partial_result_blocks_and_keeps_original_pending_view(workspace):
    report = workbook(workspace, [
        row(**{"Local User Action": "Remove only: alice; bob" + EXCLUSIONS}),
    ])
    output = workspace / "email.txt"
    update_report(report, results=[evidence(workspace, stopped())], email_output=output)
    book = load_workbook(report)
    target = records(book["Azure Data"])[0]
    assert target["Remediation Status"] == "Blocked"
    assert "alice" in target["Local User Action"]
    assert "Removed: alice; bob" not in target["Local User Action"]
    assert target["Local User Action"].endswith(EXCLUSIONS)
    assert target["Blocker"]
    assert target["Next Action"]
    pending = records(book["Pending Remediation"])
    assert len(pending) == 1
    assert pending[0]["Remediation Status"] == "Blocked"
    assert pending[0]["Blocker"] == target["Blocker"]
    email = output.read_text(encoding="utf-8")
    assert "alice" in email and "bob" in email
    assert "blocked" in email.casefold()
    book.close()


@pytest.mark.parametrize("status", ["Pending remediation", "Remediated"])
def test_absent_result_preserves_status_without_claiming_new_verified_removal(workspace, status):
    action = (
        "Removed: alice (Run Command verified)." if status == "Remediated" else "Remove only: alice"
    ) + EXCLUSIONS
    report = workbook(workspace, [
        row(**{"Remediation Status": status, "Local User Action": action}),
    ])
    update_report(
        report, results=[evidence(workspace, absent())], email_output=workspace / "email.txt",
    )
    book = load_workbook(report)
    target = records(book["Azure Data"])[0]
    assert target["Remediation Status"] == status
    assert target["Local User Action"] == action
    book.close()


@pytest.mark.parametrize("view,status", [
    ("Pending Review", "Pending review"),
    ("Pending Remediation", "Pending remediation"),
])
def test_manual_blocker_preserves_pending_view_and_can_be_cleared(workspace, view, status):
    report = workbook(workspace, [row(**{"Remediation Status": status})])
    updates = manual(workspace, [blocker()])
    update_report(report, blockers=updates, email_output=workspace / "blocked.txt")
    book = load_workbook(report)
    target = records(book["Azure Data"])[0]
    assert target["Remediation Status"] == "Blocked"
    assert target["Blocker"] == "Owner validation outstanding"
    assert len(records(book[view])) == 1
    assert records(book[view])[0]["Remediation Status"] == "Blocked"
    book.close()
    text = (workspace / "blocked.txt").read_text(encoding="utf-8")
    assert "Owner validation outstanding" in text
    assert "Obtain application owner validation" in text
    updates = manual(workspace, [
        blocker(Status=status, Blocker="", **{"Next Action": "Retry approved execution"}),
    ])
    update_report(report, blockers=updates, email_output=workspace / "cleared.txt")
    book = load_workbook(report)
    target = records(book["Azure Data"])[0]
    assert target["Remediation Status"] == status
    assert target["Blocker"] in ("", None)
    assert target["Next Action"] == "Retry approved execution"
    book.close()


@pytest.mark.parametrize("literal", ["=1+1", "+SUM(1,1)", "-1+2", "@SUM(1,1)"])
def test_manual_formula_looking_values_remain_literal_excel_strings(workspace, literal):
    report = workbook(workspace)
    updates = manual(workspace, [blocker(Blocker=literal, **{"Next Action": literal})])
    update_report(report, blockers=updates, email_output=workspace / "email.txt")
    book = load_workbook(report)
    for name in ("Azure Data", "Pending Remediation"):
        headers = [cell.value for cell in book[name][1]]
        for field in ("Blocker", "Next Action"):
            cell = book[name].cell(2, headers.index(field) + 1)
            assert cell.value == literal
            assert cell.data_type == "s"
    book.close()


@pytest.mark.parametrize("update", [
    blocker(Blocker=""),
    blocker(**{"Next Action": ""}),
    blocker(Status="Remediated"),
    blocker(Status="Pending remediation", Blocker="", **{"Next Action": ""}),
    blocker(Status="Pending review", Blocker="", **{"Next Action": ""}),
    blocker(**{"Resource Name": "unknown-vm"}),
])
def test_invalid_manual_update_fails_without_writes(workspace, update):
    report = workbook(workspace)
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(report, blockers=manual(workspace, [update]), email_output=output)
    assert_unchanged(report, original, output)


def test_result_and_manual_update_for_same_resource_rejected_atomically(workspace):
    report = workbook(workspace)
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(
            report, results=[evidence(workspace)], blockers=manual(workspace, [blocker()]),
            email_output=output,
        )
    assert_unchanged(report, original, output)


@pytest.mark.parametrize("text", [
    success(vm="unknown-vm"),
    success(users=("unapproved-user",)),
    success(users=("root",)),
    success(users=("alice", "bob")),
    success().replace(FINAL, ""),
])
def test_invalid_result_scope_fails_before_any_write(workspace, text):
    report = workbook(workspace)
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(report, results=[evidence(workspace, text)], email_output=output)
    assert_unchanged(report, original, output)


def test_valid_then_invalid_result_batch_does_not_partially_write(workspace):
    report = workbook(workspace)
    original = report.read_bytes()
    output = workspace / "email.txt"
    sources = [
        evidence(workspace),
        evidence(workspace, success(vm="unknown-vm"), name="unknown.txt"),
    ]
    with pytest.raises(ValueError):
        update_report(report, results=sources, email_output=output)
    assert_unchanged(report, original, output)


def test_success_must_cover_all_approved_users(workspace):
    report = workbook(workspace, [row(**{"Local User Action": "Remove only: alice; bob"})])
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(report, results=[evidence(workspace)], email_output=output)
    assert_unchanged(report, original, output)


def test_ambiguous_vm_requires_account_and_region_filter(workspace):
    report = workbook(workspace, [
        row(),
        row(account="other-synthetic-account"),
        row(region="other-synthetic-region"),
    ])
    original = report.read_bytes()
    output = workspace / "email.txt"
    source = evidence(workspace)
    with pytest.raises(ValueError):
        update_report(report, results=[source], email_output=output)
    assert_unchanged(report, original, output)
    update_report(
        report, results=[source], email_output=output,
        account="synthetic-account", region="synthetic-region", ticket="SYNTHETIC-1",
    )
    book = load_workbook(report)
    assert [record["Remediation Status"] for record in records(book["Azure Data"])] == [
        "Remediated", "Pending remediation", "Pending remediation",
    ]
    book.close()


def test_result_cannot_match_aws_only_resource(workspace):
    report = workbook(workspace, [])
    book = load_workbook(report)
    values = row()
    book["AWS Data"].append([values.get(header, "") for header in HEADERS])
    book.save(report)
    book.close()
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(report, results=[evidence(workspace)], email_output=output)
    assert_unchanged(report, original, output)


@pytest.mark.parametrize("defect", [
    "missing-pending-review", "missing-pending-remediation", "missing-header",
    "duplicate-header", "formula",
])
def test_unsupported_workbook_rejected_without_source_changes(workspace, defect):
    report = workbook(workspace)
    book = load_workbook(report)
    if defect == "missing-pending-review":
        del book["Pending Review"]
    elif defect == "missing-pending-remediation":
        del book["Pending Remediation"]
    elif defect == "missing-header":
        book["Azure Data"]["A1"] = "Wrong Header"
    elif defect == "duplicate-header":
        book["Azure Data"]["F1"] = "Account"
    else:
        book["Cover"]["B1"] = "=1+1"
    book.save(report)
    book.close()
    original = report.read_bytes()
    output = workspace / "email.txt"
    with pytest.raises(ValueError):
        update_report(report, results=[evidence(workspace)], email_output=output)
    assert_unchanged(report, original, output)


def test_existing_email_is_not_overwritten_and_report_is_unchanged(workspace):
    report = workbook(workspace)
    original = report.read_bytes()
    output = workspace / "email.txt"
    output.write_text("Existing synthetic draft", encoding="utf-8")
    with pytest.raises((ValueError, FileExistsError)):
        update_report(report, results=[evidence(workspace)], email_output=output)
    assert report.read_bytes() == original
    assert output.read_text(encoding="utf-8") == "Existing synthetic draft"


def test_refuses_email_destination_inside_repository(workspace):
    report = workbook(workspace)
    original = report.read_bytes()
    output = REPOSITORY / f".synthetic-email-{uuid4().hex}.txt"
    try:
        with pytest.raises(ValueError):
            update_report(report, results=[evidence(workspace)], email_output=output)
        assert_unchanged(report, original, output)
    finally:
        output.unlink(missing_ok=True)


def test_refuses_report_inside_repository(workspace):
    source = workbook(workspace)
    report = REPOSITORY / f".synthetic-report-{uuid4().hex}.xlsx"
    report.write_bytes(source.read_bytes())
    original = report.read_bytes()
    output = workspace / "email.txt"
    try:
        with pytest.raises(ValueError):
            update_report(report, results=[evidence(workspace)], email_output=output)
        assert_unchanged(report, original, output)
    finally:
        report.unlink(missing_ok=True)


def test_fully_collapsed_multi_account_output_is_parsed():
    text = success(users=("alice", "bob")).replace("\n", "")
    result = parse_result("[stdout]" + text + "[stderr]")
    assert result.verified == ("alice", "bob")


def test_known_inactive_export_note_does_not_mask_success():
    text = success().replace(
        "PRECHECKS=",
        "WARNING: Earlier Wiz export marked this VM inactive; "
        "confirm current ownership/state.\nPRECHECKS=",
    )
    assert parse_result(text).verified == ("alice",)


def test_stop_before_account_listing_is_recorded_as_blocker(workspace):
    report = workbook(workspace)
    text = (
        "VM=synthetic-vm EXECUTION=SINGLE_RUN\n"
        "STOP=Required dependency inspection unavailable\n"
        "REMOVAL_COMMANDS_SUCCEEDED=none\nRESULT=STOPPED_REVIEW_REQUIRED"
    )
    result = update_report(
        report, results=[evidence(workspace, text)], email_output=workspace / "email.txt",
    )
    book = load_workbook(report)
    assert records(book["Azure Data"])[0]["Remediation Status"] == "Blocked"
    assert json.loads(result.memory_path.read_text())["events"][0]["details"]["verified"] == []
    assert "Required dependency inspection unavailable" in (workspace / "email.txt").read_text()


def test_block_clear_block_again_is_not_deduplicated(workspace):
    report = workbook(workspace)
    first = manual(workspace, [blocker()])
    update_report(report, blockers=first, email_output=workspace / "one.txt")
    clear = manual(workspace, [
        blocker(
            Status="Pending remediation", Blocker="", **{"Next Action": "Owner approved retry"},
        ),
    ])
    update_report(report, blockers=clear, email_output=workspace / "two.txt")
    again = manual(workspace, [blocker()])
    result = update_report(report, blockers=again, email_output=workspace / "three.txt")
    assert result.imported == 1
    assert records(load_workbook(report)["Azure Data"])[0]["Remediation Status"] == "Blocked"


def test_success_after_manual_blocker_can_use_original_approval(workspace):
    report = workbook(workspace)
    update_report(
        report, blockers=manual(workspace, [blocker()]), email_output=workspace / "blocked.txt",
    )
    update_report(report, results=[evidence(workspace)], email_output=workspace / "completed.txt")
    values = records(load_workbook(report)["Azure Data"])[0]
    assert values["Remediation Status"] == "Remediated"
    assert not values["Blocker"] and not values["Next Action"]


def test_partial_then_remaining_success_accumulates_verified_scope(workspace):
    report = workbook(workspace, [row(**{"Local User Action": "Remove only: alice; bob"})])
    update_report(report, results=[evidence(workspace, stopped())],
                  email_output=workspace / "partial.txt")
    second = success(users=("bob",)).replace(
        "ACCOUNT=bob", "ACCOUNT=alice ALREADY_ABSENT (not proof of who removed it)\nACCOUNT=bob",
    )
    result = update_report(
        report, results=[evidence(workspace, second)], email_output=workspace / "completed.txt",
    )
    values = records(load_workbook(report)["Azure Data"])[0]
    assert values["Remediation Status"] == "Remediated"
    assert "alice; bob" in values["Local User Action"]
    assert json.loads(result.memory_path.read_text())["events"][-1]["details"]["verified"] == [
        "alice", "bob",
    ]


def test_success_with_previously_absent_unproven_user_is_not_full_completion(workspace):
    report = workbook(workspace, [row(**{"Local User Action": "Remove only: alice; bob"})])
    text = success().replace(
        "ACCOUNT=alice", "ACCOUNT=bob ALREADY_ABSENT (not proof of who removed it)\nACCOUNT=alice",
    )
    update_report(report, results=[evidence(workspace, text)], email_output=workspace / "draft.txt")
    values = records(load_workbook(report)["Azure Data"])[0]
    assert values["Remediation Status"] == "Pending remediation"
    assert "Reconcile" in values["Next Action"]
    assert "removed alice (partial" in (workspace / "draft.txt").read_text()


def test_raw_logs_stay_out_of_excel_cells_and_comments(workspace):
    report = workbook(workspace)
    result = update_report(
        report, results=[evidence(workspace)], email_output=workspace / "draft.txt",
    )
    book = load_workbook(report)
    for sheet in book:
        for cells in sheet:
            for cell in cells:
                for value in (str(cell.value), cell.comment.text if cell.comment else ""):
                    assert "BACKUP_PATH" not in value
                    assert "UID=" not in value
                    assert "PRECHECKS=" not in value
    assert success() in [
        e["raw_output"] for e in json.loads(result.memory_path.read_text())["events"]
    ]


def test_corrupt_memory_fails_without_changing_workbook(workspace):
    report = workbook(workspace)
    memory = report.with_suffix(".remediation-memory.json")
    memory.write_text('{"version":1,"events":[{}]}')
    original = report.read_bytes()
    with pytest.raises(ValueError, match="memory"):
        update_report(report, results=[evidence(workspace)], email_output=workspace / "email.txt")
    assert report.read_bytes() == original


def test_scope_change_requires_reconciliation(workspace):
    report = workbook(workspace)
    update_report(report, results=[evidence(workspace)], email_output=workspace / "first.txt")
    book = load_workbook(report)
    book["Azure Data"].cell(2, HEADERS.index("Local User Action") + 1).value = (
        "Removed: bob (Run Command verified)."
    )
    book.save(report)
    original = report.read_bytes()
    with pytest.raises(ValueError, match="scope"):
        update_report(report, results=[evidence(workspace)], email_output=workspace / "next.txt")
    assert report.read_bytes() == original


@pytest.mark.parametrize("existing_memory", [False, True])
def test_publish_failure_restores_memory_and_removes_unsent_draft(
    workspace, monkeypatch, existing_memory,
):
    from remediation import execution_update

    report = workbook(workspace)
    if existing_memory:
        update_report(
            report, blockers=manual(workspace, [blocker()]), email_output=workspace / "prior.txt",
        )
    memory = report.with_suffix(".remediation-memory.json")
    previous = memory.read_bytes() if memory.exists() else None
    original = report.read_bytes()
    publish = execution_update._publish

    def fail_report(stream, path, **kwargs):
        if path == report:
            raise PermissionError("Synthetic locked workbook")
        return publish(stream, path, **kwargs)

    monkeypatch.setattr(execution_update, "_publish", fail_report)
    with pytest.raises(PermissionError):
        update_report(report, results=[evidence(workspace)], email_output=workspace / "draft.txt")
    assert report.read_bytes() == original
    assert (memory.read_bytes() if memory.exists() else None) == previous
    assert not (workspace / "draft.txt").exists()


def test_cli_routes_existing_csv_and_new_update_command(workspace, monkeypatch, capsys):
    from remediation import __main__, execution_update

    calls = []
    monkeypatch.setattr(__main__, "report_main", lambda args: calls.append(args) or 0)
    assert __main__.main(["synthetic.csv"]) == 0
    assert calls == [["synthetic.csv"]]
    report = workbook(workspace)
    result = evidence(workspace)
    assert __main__.main([
        "update-remediation", "--report", str(report), "--result", str(result),
        "--email-output", str(workspace / "email.txt"),
    ]) == 0
    assert "Imported: 1" in capsys.readouterr().out
    assert execution_update.main([
        "--report", str(report), "--email-output", str(workspace / "nothing.txt"),
    ]) == 1
    assert "Update failed" in capsys.readouterr().err


def test_email_only_preserves_report_and_does_not_create_memory(workspace):
    report = workbook(workspace, [row(**{
        "Remediation Status": "Remediated",
        "Local User Action": "Removed: alice (Run Command verified)." + EXCLUSIONS,
    })])
    original = report.read_bytes()
    result = update_report(
        report, email_only=True, email_output=workspace / "draft.txt", ticket="CHANGE-EXAMPLE",
    )
    assert report.read_bytes() == original
    assert result.backup_path is None
    assert not result.memory_path.exists()
    draft = (workspace / "draft.txt").read_text()
    assert "removed alice" in draft
    assert "CHANGE-EXAMPLE" in draft
    assert "Excluded" in draft
    with pytest.raises(ValueError):
        update_report(
            report, email_only=True, results=[evidence(workspace)],
            email_output=workspace / "invalid.txt",
        )


def test_overlong_visible_blocker_fails_without_clipping_or_writing(workspace):
    report = workbook(workspace)
    original = report.read_bytes()
    with pytest.raises(ValueError, match="height"):
        update_report(
            report, blockers=manual(workspace, [blocker(Blocker="detail " * 400)]),
            email_output=workspace / "draft.txt",
        )
    assert report.read_bytes() == original
    assert not (workspace / "draft.txt").exists()
